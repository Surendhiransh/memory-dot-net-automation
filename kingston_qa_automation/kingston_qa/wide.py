"""Shared row-wise report layout used by every output (servers, memory, SSD, USB).

Purpose : one line per record. Each field gets three columns: CSV value | Website value | Status.
          Writes the same table as:
            - CSV  : flat headers such as "Part Number (CSV)"
            - Excel: 'Summary' sheet (live COUNTIF formulas) + 'Results' sheet (two-row header,
                     light-blue header, colour-coded statuses, frozen header, filters)
            - HTML : a static report you can open in any browser
Used by : reporter.py (servers) and parts_reporter.py (memory / SSD / USB).
"""
from __future__ import annotations

import csv
import html
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEADER_FILL, SUBHEADER_FILL, HEADER_TEXT, BORDER_COLOR = "DDEBF7", "EEF5FB", "1F3864", "BFBFBF"
MAX_CELL = 32000  # Excel limit is 32,767 characters per cell
ROW_HEIGHT = 45   # fixed so long JSON cells stay tidy; the full text is in the cell
PALETTE = {  # status -> (light fill, dark text)
    "MATCHED": ("E2EFDA", "375623"), "MISMATCHED": ("FCE4E4", "9C0006"), "UNMATCHED": ("FFF2CC", "7F6000"),
    "BLOCKED": ("FBE5D6", "843C0C"), "SKIPPED": ("EDEDED", "595959"),
    "Matched": ("E2EFDA", "375623"), "Mismatched": ("FCE4E4", "9C0006"), "Unmatched": ("FFF2CC", "7F6000"),
    "PASS": ("E2EFDA", "375623"), "FAIL": ("FCE4E4", "9C0006"), "NEEDS REVIEW": ("EDE7F6", "5B2C83"),
}

_THIN = Side(style="thin", color=BORDER_COLOR)
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BODY = Font(name=FONT, size=10)
_NOTE = Font(name=FONT, italic=True, size=9, color="595959")
_WRAP = Alignment(vertical="top", wrap_text=True)
_CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


@dataclass
class WideTable:
    title: str
    lead: list[str]                  # identity columns, e.g. ["Row No", "Overall Status"]
    groups: list[str]                # one title per validated field
    parts: list[str]                 # sub-columns per field, e.g. ["CSV", "Website", "Status"]
    tail: list[str]                  # trailing columns, e.g. ["Comments"]
    rows: list[list[Any]]            # lead + groups*parts + tail values per record
    overall_col: int                 # 1-based position of the overall status within `lead`
    overall_values: list[str]
    field_values: list[str]          # values that appear in each group's status sub-column
    status_part: int                 # 0-based index of the status sub-column within `parts`
    legend: str = ""
    record_label: str = "Records"
    wide_groups: set[str] = field(default_factory=set)  # groups holding long text (wider columns)

    @property
    def ncols(self) -> int:
        return len(self.lead) + len(self.groups) * len(self.parts) + len(self.tail)

    def flat_headers(self) -> list[str]:
        return self.lead + [f"{g} ({p})" for g in self.groups for p in self.parts] + self.tail

    def status_cols(self) -> list[int]:
        start = len(self.lead) + 1 + self.status_part
        return [self.overall_col] + [start + i * len(self.parts) for i in range(len(self.groups))]


def _xl(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return value if len(value) <= MAX_CELL else value[:MAX_CELL] + " ...[truncated]"


# ---------------------------------------------------------------- CSV
def write_csv(table: WideTable, path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(table.flat_headers())
        writer.writerows(table.rows)


# ---------------------------------------------------------------- Excel
def _header(ws, row: int, first: int, last: int, fill: str = HEADER_FILL) -> None:
    ws.row_dimensions[row].height = 24
    for col in range(first, last + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT, bold=True, color=HEADER_TEXT, size=11)
        cell.fill = PatternFill("solid", start_color=fill)
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _status(cell) -> None:
    fill, text = PALETTE.get(str(cell.value), ("FFFFFF", "000000"))
    cell.fill = PatternFill("solid", start_color=fill)
    cell.font = Font(name=FONT, size=10, bold=True, color=text)
    cell.alignment = _CENTER


def _print_setup(ws) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _results_sheet(ws, t: WideTable) -> int:
    ws.title = "Results"
    lead, parts = len(t.lead), len(t.parts)
    tail_start = lead + len(t.groups) * parts + 1
    ws.append(t.lead + [h for g in t.groups for h in [g] + [""] * (parts - 1)] + t.tail)
    ws.append([""] * lead + t.parts * len(t.groups) + [""] * len(t.tail))
    _header(ws, 1, 1, t.ncols)
    _header(ws, 2, 1, t.ncols)
    for col in range(lead + 1, tail_start):
        ws.cell(row=2, column=col).fill = PatternFill("solid", start_color=SUBHEADER_FILL)
    for col in list(range(1, lead + 1)) + list(range(tail_start, t.ncols + 1)):
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    for i in range(len(t.groups)):
        start = lead + i * parts + 1
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=start + parts - 1)

    status_cols = t.status_cols()
    for values in t.rows:
        ws.append([_xl(v) for v in values])
        r = ws.max_row
        ws.row_dimensions[r].height = ROW_HEIGHT
        for col in range(1, t.ncols + 1):
            cell = ws.cell(row=r, column=col)
            cell.font, cell.border, cell.alignment = _BODY, _BORDER, _WRAP
        ws.cell(row=r, column=1).alignment = _CENTER
        for col in status_cols:
            _status(ws.cell(row=r, column=col))

    widths = [9 if i == 0 else 16 for i in range(lead)]
    for g in t.groups:
        value_width = 50 if g in t.wide_groups else 28
        widths += [13 if i == t.status_part else value_width for i in range(parts)]
    widths += [70] * len(t.tail)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=3, column=lead + 1)
    last = max(ws.max_row, 3)
    ws.auto_filter.ref = f"A2:{get_column_letter(t.ncols)}{last}"
    ws.sheet_view.zoomScale = 90
    ws.print_title_rows = "1:2"
    _print_setup(ws)
    return last


def _count_block(ws, headers: list[str], labels: list[str], formulas) -> None:
    ws.append([])
    ws.append(headers)
    _header(ws, ws.max_row, 1, len(headers))
    first = ws.max_row + 1
    for label in labels:
        r = ws.max_row + 1
        ws.append([label, *formulas(r)])
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=col)
            cell.font, cell.border, cell.alignment = _BODY, _BORDER, _CENTER
        _status(ws.cell(row=r, column=1))
    end = ws.max_row
    ws.append(["Total", *[f"=SUM({get_column_letter(c)}{first}:{get_column_letter(c)}{end})"
                          for c in range(2, len(headers) + 1)]])
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=ws.max_row, column=col)
        cell.font = Font(name=FONT, size=10, bold=True)
        cell.fill = PatternFill("solid", start_color="F2F2F2")
        cell.border, cell.alignment = _BORDER, _CENTER


def _summary_sheet(ws, t: WideTable, last: int, meta: dict[str, Any]) -> None:
    ws["A1"] = f"{t.title} - Validation Summary"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=HEADER_TEXT)
    ws["A2"] = f"Generated {datetime.now():%Y-%m-%d %H:%M:%S}"
    ws["A2"].font = _NOTE
    overall = get_column_letter(t.overall_col)
    overall_rng = f"Results!{overall}3:{overall}{last}"
    data_rng = f"Results!A3:{get_column_letter(t.ncols)}{last}"
    _count_block(ws, ["Overall Status", t.record_label], t.overall_values,
                 lambda r: [f"=COUNTIF({overall_rng},A{r})"])
    # Field statuses only appear in status cells; subtract the overall column in case vocabularies overlap.
    _count_block(ws, ["Field Status", "Field Checks"], t.field_values,
                 lambda r: [f"=COUNTIF({data_rng},A{r})-COUNTIF({overall_rng},A{r})"])
    if t.legend:
        ws.append([t.legend])
        ws.cell(row=ws.max_row, column=1).font = _NOTE
    ws.append([])
    ws.append(["Run Detail", "Value"])
    _header(ws, ws.max_row, 1, 3)
    ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=3)
    for key, value in meta.items():
        ws.append([str(key), _xl(str(value))])
        for col in (1, 2, 3):
            cell = ws.cell(row=ws.max_row, column=col)
            cell.font, cell.border, cell.alignment = _BODY, _BORDER, _WRAP
        ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=3)
    for col, width in zip("ABC", (24, 18, 60)):
        ws.column_dimensions[col].width = width
    ws.sheet_view.showGridLines = False
    _print_setup(ws)


def write_xlsx(table: WideTable, path: Path, meta: dict[str, Any]) -> None:
    wb = Workbook()
    last = _results_sheet(wb.active, table)
    _summary_sheet(wb.create_sheet("Summary", 0), table, last, meta)
    wb.save(path)


# ---------------------------------------------------------------- HTML
_HTML_COLORS = {"MATCHED": "#1a7f37", "Matched": "#1a7f37", "PASS": "#1a7f37", "MISMATCHED": "#cf222e",
                "Mismatched": "#cf222e", "FAIL": "#cf222e", "UNMATCHED": "#9a6700", "Unmatched": "#9a6700",
                "BLOCKED": "#bc4c00", "SKIPPED": "#57606a", "NEEDS REVIEW": "#8250df"}


def write_html(table: WideTable, path: Path, meta: dict[str, Any], counts: dict[str, int]) -> None:
    e = html.escape
    status_cols = {c - 1 for c in table.status_cols()}

    def cell(i: int, v: Any) -> str:
        text = "" if v is None else str(v)
        if i in status_cols and text:
            return f'<td><span class="b" style="background:{_HTML_COLORS.get(text, "#333")}">{e(text)}</span></td>'
        return f"<td><div class=\"v\">{e(text)}</div></td>"

    tiles = "".join(f'<div class="t"><div class="n" style="color:{_HTML_COLORS.get(s, "#333")}">{n}</div>'
                    f"<div>{e(s)}</div></div>" for s, n in counts.items())
    head1 = ("".join(f'<th rowspan="2">{e(h)}</th>' for h in table.lead)
             + "".join(f'<th colspan="{len(table.parts)}">{e(g)}</th>' for g in table.groups)
             + "".join(f'<th rowspan="2">{e(h)}</th>' for h in table.tail))
    head2 = "".join(f'<th class="sub">{e(p)}</th>' for _ in table.groups for p in table.parts)
    body = "".join("<tr>" + "".join(cell(i, v) for i, v in enumerate(r)) + "</tr>" for r in table.rows)
    meta_rows = "".join(f"<tr><th>{e(str(k))}</th><td>{e(str(v))}</td></tr>" for k, v in meta.items())
    path.write_text(f"""<!doctype html><html><head><meta charset="utf-8"><title>{e(table.title)}</title>
<style>
body{{font:13px Arial,sans-serif;margin:20px;color:#1f2328;background:#fff}}
table{{border-collapse:collapse;margin:8px 0}}th,td{{border:1px solid #d0d7de;padding:5px;vertical-align:top;text-align:left}}
th{{background:#ddebf7;color:#1f3864;text-align:center}}th.sub{{background:#eef5fb}}
.v{{max-width:320px;max-height:120px;overflow:auto;white-space:pre-wrap;word-break:break-word}}
.b{{color:#fff;border-radius:4px;padding:2px 6px;font-size:11px;white-space:nowrap}}
.tiles{{display:flex;gap:12px;flex-wrap:wrap}}.t{{border:1px solid #d0d7de;border-radius:6px;padding:10px 16px;text-align:center}}
.n{{font-size:24px;font-weight:700}}.wrap{{overflow-x:auto}}
</style></head><body><h1>{e(table.title)} Report</h1>
<p>Generated {datetime.now():%Y-%m-%d %H:%M:%S} &middot; {len(table.rows)} {e(table.record_label.lower())}</p>
<div class="tiles">{tiles}</div><p><small>{e(table.legend)}</small></p>
<h2>Run</h2><table>{meta_rows}</table>
<h2>Results</h2><div class="wrap"><table><tr>{head1}</tr><tr>{head2}</tr>{body}</table></div></body></html>""",
                    encoding="utf-8")
