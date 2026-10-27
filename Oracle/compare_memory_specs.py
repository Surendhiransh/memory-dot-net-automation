"""
Compare memory specifications in oracle_host_memory_specification.csv against the
source Oracle / Fujitsu datasheet PDFs.

Fields compared: dimm_slots, maximum_memory, maximum_LRDIMM, maximum_RDIMM,
maximum_UDIMM, maximum_SODIMM.

Rules
-----
* dimm_slots      : Match if CSV equals the total slot count OR the per-processor count.
* maximum_memory  : Spec-table value ("... maximum memory configuration", "DIMM slots provide
                    up to ...") wins over summary/marketing text. If CSV matches the spec table it is
                    a Match; a disagreeing summary is only noted in the comment.
* DIMM type sizes : Largest size listed for that DIMM type (e.g. "RDIMM options: 16 GB and 32 GB" -> 32 GB).
* Missing values  : 'nan' in CSV and nothing in PDF -> Match. Present on only one side -> Mismatch.
* Missing columns : a CSV without one of the required columns is rejected before comparing.

Usage
-----
    python compare_memory_specs.py [--csv PATH] [--pdf-dir PATH] [--out-dir PATH]

Without --csv, oracle_host_memory_specification.csv is used; if it does not exist, the newest
browser-download copy such as "oracle_host_memory_specification (1).csv" is used instead.
"""

import argparse
import ast
import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent

FIELDS = [
    "dimm_slots",
    "maximum_memory",
    "maximum_LRDIMM",
    "maximum_RDIMM",
    "maximum_UDIMM",
    "maximum_SODIMM",
]

REQUIRED_COLUMNS = FIELDS + ["file_name"]
DEFAULT_CSV_NAME = "oracle_host_memory_specification.csv"

MATCH, MISMATCH = "Match", "Mismatch"

# --------------------------------------------------------------------------- number / size helpers

_UNITS = {
    "zero": 0, "one": 1, "single": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6,
    "seven": 7, "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12, "thirteen": 13,
    "fourteen": 14, "fifteen": 15, "sixteen": 16, "seventeen": 17, "eighteen": 18, "nineteen": 19,
}
_TENS = {
    "twenty": 20, "thirty": 30, "forty": 40, "fifty": 50, "sixty": 60,
    "seventy": 70, "eighty": 80, "ninety": 90,
}
NUM_WORD = (
    r"(?:\d[\d,]*|(?:" + "|".join(_TENS) + r")(?:-(?:" + "|".join(k for k in _UNITS if k != "single")
    + r"))?|" + "|".join(_UNITS) + r")"
)
SIZE = r"(\d+(?:\.\d+)?)\s*(PB|TB|GB)\b"
_UNIT_TO_GB = {"GB": 1, "TB": 1024, "PB": 1024 * 1024}


def word_to_int(text):
    text = text.strip().lower().replace(",", "")
    if text.isdigit():
        return int(text)
    if text in _UNITS:
        return _UNITS[text]
    tens, _, unit = text.partition("-")
    return _TENS[tens] + (_UNITS[unit] if unit else 0)


def size_to_gb(value, unit):
    return float(value) * _UNIT_TO_GB[unit.upper()]


def format_size(gb):
    """Render a GB amount the way the datasheets do (768 GB, 1.5 TB, 6 TB)."""
    if gb is None:
        return ""
    if gb >= 1024:
        tb = gb / 1024
        return f"{tb:g} TB"
    return f"{gb:g} GB"


def parse_csv_size(text):
    m = re.fullmatch(r"\s*" + SIZE + r"\s*", text, re.I)
    if not m:
        raise ValueError(f"Unrecognised size value in CSV: {text!r}")
    return size_to_gb(m.group(1), m.group(2))


def parse_csv_cell(raw):
    """CSV cells look like "['24']", "['1.5 TB']" or "['nan']". Returns a list of non-empty strings."""
    raw = (raw or "").strip()
    try:
        value = ast.literal_eval(raw)
    except (ValueError, SyntaxError):
        value = raw
    items = value if isinstance(value, (list, tuple)) else [value]
    cleaned = [str(v).strip() for v in items]
    return [v for v in cleaned if v and v.lower() not in ("nan", "none", "null", "n/a", "-")]


# --------------------------------------------------------------------------- PDF extraction

@dataclass
class Finding:
    value: float          # slot count or size in GB
    page: int
    snippet: str
    kind: str = ""        # e.g. "spec", "summary", "total", "per_processor"


@dataclass
class PdfSpecs:
    pages: list = field(default_factory=list)
    slots_total: list = field(default_factory=list)
    slots_per_proc: list = field(default_factory=list)
    max_processors: list = field(default_factory=list)
    memory_spec: list = field(default_factory=list)
    memory_summary: list = field(default_factory=list)
    dimm: dict = field(default_factory=lambda: {k: [] for k in ("LRDIMM", "RDIMM", "UDIMM", "SODIMM")})
    generic_dimm: list = field(default_factory=list)


def normalise(text):
    text = text.replace("­", "").replace(" ", " ")
    text = re.sub(r"[•●▪�]", " • ", text)   # bullet glyphs -> clause separator
    text = text.replace("’", "'").replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", text)


def snippet(text, m, width=70):
    start, end = max(0, m.start() - width), min(len(text), m.end() + width)
    return ("..." if start else "") + text[start:end].strip() + ("..." if end < len(text) else "")


SLOT_NOUN = (
    r"(?:DDR\d(?:-\d+)?\s+)?"
    r"(?:dual inline memory modules?\s*\(DIMMs?\)\s*slots|memory modules?\s*\(DIMMs?\)\s*slots"
    r"|DIMM slots|DIMMs)"
)
RE_SLOTS = re.compile(
    r"(?<![-\w])(" + NUM_WORD + r")\s+" + SLOT_NOUN + r"(?!\s*per channel)(\s+per (?:processor|socket|CPU))?",
    re.I,
)
RE_SLOTS_PAREN = re.compile(  # "Up to 192 DIMMs (24 per CPU module)", "Forty-eight DIMM slots (12 per CPU module)"
    r"\((\d+)\s+(?:DIMMs\s+)?per (?:processor|socket|CPU module|CPU)\)", re.I
)
RE_PROCS = [
    re.compile(r"\b(" + NUM_WORD + r")-processor (?:system|server)", re.I),
    re.compile(r"\b(single) processor\b", re.I),
    re.compile(r"configured with (?:one or )?(" + NUM_WORD + r")\b[^.•]{0,30}?\bprocessors?\b(?!\s*cores?)", re.I),
    re.compile(r"\bup to (" + NUM_WORD + r") (?:processors\b|CPUs\b)(?!\s*cores?)", re.I),
]
RE_MEM_SPEC = [
    re.compile(SIZE + r" maximum memory (?:configuration|capacity)", re.I),
    re.compile(r"\bMaximum memory capacity of " + SIZE, re.I),
    re.compile(r"DIMM slots(?: \([^)]*\))? provide up to " + SIZE, re.I),
    re.compile(r"Up to " + SIZE + r" per unit, with \d+ GB DIMM", re.I),
]
RE_MEM_SUMMARY = [
    re.compile(r"slots,? with maximum memory(?: capacity)? of " + SIZE, re.I),
    re.compile(r"\bup to " + SIZE + r" of (?:main )?memory(?! in a (?:single )?rack)", re.I),
    re.compile(r"memory capacity (?:of )?up to " + SIZE, re.I),
]
# "<TYPE> options: 16 GB at DDR4-2666 and 32 GB at ..." - the clause is cut at the next bullet,
# DIMM type keyword or unrelated text (some PDFs run the option line straight into sidebar text).
RE_DIMM = {
    "LRDIMM": [
        re.compile(r"(?:LRDIMM|Load-reduced DIMM) options?:\s*([^•]{0,160})", re.I),
        re.compile(r"\b(\d+) GB\s+(?:[\w.,-]+\s+){0,5}?(?:LRDIMMs?\b|load-reduced\s+(?:ECC\s+)?DIMMs?)", re.I),
    ],
    "RDIMM": [
        re.compile(r"(?<![\w-])RDIMM options?:\s*([^•]{0,160})", re.I),
        re.compile(r"\b(\d+) GB\s+(?:[\w.,-]+\s+){0,5}?registered\s+(?:low-voltage\s+)?(?:ECC\s+)?DIMMs?\b", re.I),
    ],
    "UDIMM": [
        re.compile(r"(?<![\w-])UDIMM options?:\s*([^•]{0,160})", re.I),
        re.compile(r"\b(\d+) GB\s+(?:[\w.,-]+\s+){0,5}?(?:UDIMMs?\b|unbuffered\s+(?:ECC\s+)?DIMMs?)", re.I),
    ],
    "SODIMM": [
        re.compile(r"\bSO-?DIMM options?:\s*([^•]{0,160})", re.I),
        re.compile(r"\b(\d+) GB\s+(?:[\w.,-]+\s+){0,5}?SO-?DIMMs?\b", re.I),
    ],
}
RE_CLAUSE_END = re.compile(r"•|L?RDIMM|[UR]DIMM|SO-?DIMM|Load-reduced|Oracle|\.\s", re.I)
RE_DIMM_SIZE = re.compile(r"(\d+)\s*GB\b", re.I)
RE_GENERIC_DIMM = re.compile(r"with (\d+) GB DIMMs?\b", re.I)


def extract_slots(text, page_no, specs):
    for match in RE_SLOTS.finditer(text):
        count = word_to_int(match.group(1))
        if count < 2:  # "one DIMM ..." style prose, not a slot count
            continue
        kind = "per_processor" if match.group(2) else "total"
        target = specs.slots_per_proc if match.group(2) else specs.slots_total
        target.append(Finding(count, page_no, snippet(text, match), kind))

    for match in RE_SLOTS_PAREN.finditer(text):
        specs.slots_per_proc.append(
            Finding(int(match.group(1)), page_no, snippet(text, match), "per_processor")
        )


def extract_processors(text, page_no, specs):
    for regex in RE_PROCS:
        for match in regex.finditer(text):
            specs.max_processors.append(
                Finding(word_to_int(match.group(1)), page_no, snippet(text, match))
            )


def extract_memory(text, page_no, specs):
    for regex in RE_MEM_SPEC:
        for match in regex.finditer(text):
            specs.memory_spec.append(
                Finding(size_to_gb(match.group(1), match.group(2)), page_no, snippet(text, match), "spec")
            )

    for regex in RE_MEM_SUMMARY:
        for match in regex.finditer(text):
            specs.memory_summary.append(
                Finding(size_to_gb(match.group(1), match.group(2)), page_no, snippet(text, match), "summary")
            )


def extract_dimms(text, page_no, specs):
    for dimm_type, (options_regex, inline_regex) in RE_DIMM.items():
        for match in options_regex.finditer(text):
            clause = RE_CLAUSE_END.split(match.group(1), maxsplit=1)[0]
            for size_match in RE_DIMM_SIZE.finditer(clause):
                specs.dimm[dimm_type].append(
                    Finding(float(size_match.group(1)), page_no, snippet(text, match))
                )
        for match in inline_regex.finditer(text):
            specs.dimm[dimm_type].append(
                Finding(float(match.group(1)), page_no, snippet(text, match))
            )

    for match in RE_GENERIC_DIMM.finditer(text):
        specs.generic_dimm.append(
            Finding(float(match.group(1)), page_no, snippet(text, match))
        )


def extract_pdf(path):
    specs = PdfSpecs()
    with pymupdf.open(path) as doc:
        specs.pages = [normalise(page.get_text()) for page in doc]

    for page_no, text in enumerate(specs.pages, start=1):
        extract_slots(text, page_no, specs)
        extract_processors(text, page_no, specs)
        extract_memory(text, page_no, specs)
        extract_dimms(text, page_no, specs)

    return specs


# --------------------------------------------------------------------------- comparison

@dataclass
class FieldResult:
    csv_value: str
    pdf_value: str
    status: str
    comment: str
    evidence: list


def distinct(values):
    return sorted({f.value for f in values})


def slot_values(specs):
    totals = distinct(specs.slots_total)
    per_processor = distinct(specs.slots_per_proc)
    processors = max((finding.value for finding in specs.max_processors), default=None)
    total = max(totals) if totals else None
    per_processor_max = max(per_processor) if per_processor else None
    derived_total = None
    if total is None and per_processor_max is not None and processors:
        derived_total = int(per_processor_max * processors)
    return totals, total, per_processor_max, processors, derived_total


def slot_display(total, per_processor, processors, derived_total):
    parts = []
    if total is not None:
        parts.append(f"{total:g} total")
    if derived_total is not None:
        parts.append(
            f"{derived_total} total (derived: {per_processor:g}/processor x {processors:g} processors)"
        )
    if per_processor is not None:
        parts.append(f"{per_processor:g} per processor")
    return "; ".join(parts)


def system_total_text(total, derived_total):
    if total is not None:
        return f"{total:g}"
    if derived_total is not None:
        return str(derived_total)
    return "unknown"


def compare_slots(csv_vals, specs):
    csv_n = int(csv_vals[0]) if csv_vals else None
    totals, total, per_proc, processors, derived_total = slot_values(specs)
    pdf_display = slot_display(total, per_proc, processors, derived_total)
    evidence = specs.slots_total + specs.slots_per_proc

    if csv_n is None and not pdf_display:
        return FieldResult("", "", MATCH, "", evidence)
    if csv_n is None:
        return FieldResult("", pdf_display, MISMATCH, f"dimm_slots missing in CSV; PDF states {pdf_display}.", evidence)
    if not pdf_display:
        return FieldResult(str(csv_n), "", MISMATCH, f"dimm_slots: CSV={csv_n} but no DIMM slot count found in PDF.", evidence)

    if total is not None and csv_n == total:
        comment = ""
        if len(totals) > 1:
            comment = f"dimm_slots: PDF mentions several totals {[int(t) for t in totals]}; CSV matches the largest."
        return FieldResult(str(csv_n), pdf_display, MATCH, comment, evidence)
    if per_proc is not None and csv_n == per_proc:
        total_txt = system_total_text(total, derived_total)
        comment = (f"dimm_slots: CSV={csv_n} matches the per-processor count; "
                   f"system total = {total_txt}.") if (total or derived_total) and (total or derived_total) != csv_n else ""
        return FieldResult(str(csv_n), pdf_display, MATCH, comment, evidence)
    if derived_total is not None and csv_n == derived_total:
        return FieldResult(str(csv_n), pdf_display, MATCH,
                           f"dimm_slots: CSV matches derived total ({per_proc:g}/processor x {processors:g}).", evidence)
    return FieldResult(str(csv_n), pdf_display, MISMATCH, f"dimm_slots: CSV={csv_n} but PDF states {pdf_display}.", evidence)


def select_memory_value(spec_values, summary_values):
    if spec_values:
        return max(spec_values), "spec table"
    if summary_values:
        return max(summary_values), "summary text (no spec-table value found)"
    return None, ""


def memory_notes(spec_values, summary_values, pdf_gb):
    notes = []
    if len(spec_values) > 1:
        others = ", ".join(format_size(value) for value in spec_values if value != pdf_gb)
        notes.append(f"PDF spec also lists smaller configuration(s): {others}")

    conflicts = [
        value
        for value in summary_values
        if pdf_gb is not None and value != pdf_gb and value not in spec_values
    ]
    if conflicts and spec_values:
        summary_sizes = ", ".join(format_size(value) for value in conflicts)
        notes.append(
            f"PDF is internally inconsistent: summary text says {summary_sizes} "
            f"but spec table says {format_size(pdf_gb)}"
        )
    return notes


def compare_memory(csv_vals, specs):
    csv_gb = parse_csv_size(csv_vals[0]) if csv_vals else None
    spec_vals = distinct(specs.memory_spec)
    summary_vals = distinct(specs.memory_summary)
    pdf_gb, source = select_memory_value(spec_vals, summary_vals)
    notes = memory_notes(spec_vals, summary_vals, pdf_gb)

    csv_display = csv_vals[0] if csv_vals else ""
    pdf_display = format_size(pdf_gb)
    evidence = specs.memory_spec + specs.memory_summary
    note_txt = ("; ".join(notes) + ".") if notes else ""

    if csv_gb is None and pdf_gb is None:
        return FieldResult("", "", MATCH, "", evidence)
    if csv_gb is None:
        return FieldResult("", pdf_display, MISMATCH,
                           f"maximum_memory missing in CSV; PDF {source} states {pdf_display}. {note_txt}".strip(), evidence)
    if pdf_gb is None:
        return FieldResult(csv_display, "", MISMATCH,
                           f"maximum_memory: CSV={csv_display} but no maximum memory found in PDF.", evidence)
    if csv_gb == pdf_gb:
        # Spec table is authoritative; a contradicting summary bullet is noted but is not an issue.
        prefix = f"maximum_memory: CSV matches PDF {source}. " if notes else ""
        return FieldResult(csv_display, pdf_display, MATCH, (prefix + note_txt).strip(), evidence)
    if csv_gb in summary_vals or csv_gb in spec_vals:
        return FieldResult(csv_display, pdf_display, MISMATCH,
                           f"maximum_memory: CSV={csv_display} appears elsewhere in the PDF, but the authoritative "
                           f"spec-table maximum is {pdf_display}. {note_txt}".strip(), evidence)
    return FieldResult(csv_display, pdf_display, MISMATCH,
                       f"maximum_memory: CSV={csv_display} but PDF {source} states {pdf_display}. {note_txt}".strip(),
                       evidence)


def compare_dimm(dimm_type, csv_vals, specs):
    csv_gb = parse_csv_size(csv_vals[0]) if csv_vals else None
    found = distinct(specs.dimm[dimm_type])
    pdf_gb = max(found) if found else None
    csv_display = csv_vals[0] if csv_vals else ""
    pdf_display = format_size(pdf_gb)
    evidence = specs.dimm[dimm_type]
    options = f" (PDF options: {', '.join(format_size(v) for v in found)})" if len(found) > 1 else ""
    name = f"maximum_{dimm_type}"

    if csv_gb is None and pdf_gb is None:
        return FieldResult("", "", MATCH, "", evidence)
    if csv_gb is None:
        return FieldResult("", pdf_display, MISMATCH, f"{name} missing in CSV; PDF states {pdf_display}{options}.", evidence)
    if pdf_gb is None:
        return FieldResult(csv_display, "", MISMATCH, f"{name}: CSV={csv_display} but PDF does not mention {dimm_type}.", evidence)
    if csv_gb == pdf_gb:
        return FieldResult(csv_display, pdf_display, MATCH, "", evidence)
    return FieldResult(csv_display, pdf_display, MISMATCH,
                       f"{name}: CSV={csv_display} but PDF largest {dimm_type} is {pdf_display}{options}.", evidence)


def compare_row(row, pdf_path):
    missing_columns = [f for f in FIELDS if f not in row]
    if missing_columns:
        return None, f"CSV row is missing column(s): {', '.join(missing_columns)}"
    if pdf_path is None:
        return None, "PDF file not found in the PDF folder."
    specs = extract_pdf(pdf_path)
    if not any(p.strip() for p in specs.pages):
        return None, "PDF has no extractable text (scanned image?) - manual check required."

    results = {
        "dimm_slots": compare_slots(parse_csv_cell(row.get("dimm_slots")), specs),
        "maximum_memory": compare_memory(parse_csv_cell(row.get("maximum_memory")), specs),
    }
    for dimm_type in ("LRDIMM", "RDIMM", "UDIMM", "SODIMM"):
        key = f"maximum_{dimm_type}"
        results[key] = compare_dimm(dimm_type, parse_csv_cell(row.get(key)), specs)

    # Informational: SPARC/Fujitsu datasheets give a DIMM size without saying RDIMM/LRDIMM.
    typed = any(specs.dimm[t] for t in specs.dimm)
    if not typed and specs.generic_dimm:
        biggest = max(f.value for f in specs.generic_dimm)
        results["maximum_LRDIMM"].comment = (results["maximum_LRDIMM"].comment + " " if results["maximum_LRDIMM"].comment else "") + \
            f"Note: PDF lists generic DDR DIMMs up to {format_size(biggest)} without stating RDIMM/LRDIMM type."
    return results, ""


def overall(results):
    statuses = [r.status for r in results.values()]
    return MISMATCH if MISMATCH in statuses else MATCH


# --------------------------------------------------------------------------- output

FILLS = {
    MATCH: PatternFill("solid", fgColor="C6EFCE"),
    MISMATCH: PatternFill("solid", fgColor="FFC7CE"),
    "Error": PatternFill("solid", fgColor="D9D9D9"),
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def result_header():
    header = ["file_name"]
    for f in FIELDS:
        header += [f"CSV_{f}", f"PDF_{f}", f"{f}_status"]
    return header + ["overall_status", "comments"]


def write_sheet(ws, rows, widths=None):
    for r in rows:
        ws.append(r)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for idx, col in enumerate(ws.columns, start=1):
        width = (widths or {}).get(idx) or min(45, max(len(str(c.value or "")) for c in col) + 2)
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.value in FILLS:
                cell.fill = FILLS[cell.value]


def safe_compare_row(row, pdf_index):
    name = (row.get("file_name") or "").strip()
    try:
        results, error = compare_row(row, pdf_index.get(name.lower()))
    except Exception as exc:  # keep processing the other files
        results, error = None, f"Error while processing: {exc}"
    return name, results, error


def error_output_row(row, name, error):
    output = [name]
    for field_name in FIELDS:
        output += [", ".join(parse_csv_cell(row.get(field_name))), "", "Error"]
    return output + ["Error", error]


def append_evidence(name, field_name, result, evidence_rows):
    seen = set()
    for finding in result.evidence:
        key = (finding.page, finding.snippet)
        if key in seen:
            continue
        seen.add(key)
        shown = (
            f"{finding.value:g} ({finding.kind})"
            if field_name == "dimm_slots"
            else format_size(finding.value)
        )
        evidence_rows.append([name, field_name, shown, finding.page, finding.snippet])


def success_output_row(name, results, counts, evidence_rows):
    output, comments = [name], []
    for field_name in FIELDS:
        result = results[field_name]
        output += [result.csv_value, result.pdf_value, result.status]
        counts[field_name][result.status] += 1
        if result.comment:
            comments.append(result.comment)
        append_evidence(name, field_name, result, evidence_rows)

    status = overall(results)
    counts["overall"][status] += 1
    comment_text = " | ".join(comments) if comments else "All fields match the PDF."
    return output + [status, comment_text]


def build_report_rows(csv_rows, pdf_index):
    result_rows = [result_header()]
    evidence_rows = [["file_name", "field", "PDF value", "page", "PDF text"]]
    counts = {field_name: {MATCH: 0, MISMATCH: 0} for field_name in FIELDS + ["overall"]}
    errors = 0

    for row in csv_rows:
        name, results, error = safe_compare_row(row, pdf_index)
        if results is None:
            errors += 1
            result_rows.append(error_output_row(row, name, error))
            continue
        result_rows.append(success_output_row(name, results, counts, evidence_rows))

    return result_rows, evidence_rows, counts, errors


def build_summary_rows(csv_rows, counts, errors):
    rows = [["field", MATCH, MISMATCH]]
    rows += [[field_name, values[MATCH], values[MISMATCH]] for field_name, values in counts.items()]
    rows += [[], ["Files processed", len(csv_rows)], ["Files with errors (missing/unreadable PDF)", errors]]
    return rows


def write_reports(out_dir, result_rows, summary_rows, evidence_rows):
    out_dir.mkdir(parents=True, exist_ok=True)
    csv_out = out_dir / "memory_spec_comparison.csv"
    xlsx_out = out_dir / "memory_spec_comparison.xlsx"

    with csv_out.open("w", newline="", encoding="utf-8-sig") as file_handle:
        csv.writer(file_handle).writerows(result_rows)

    workbook = Workbook()
    results_sheet = workbook.active
    results_sheet.title = "Results"
    write_sheet(results_sheet, result_rows, widths={1: 34, len(result_rows[0]): 90})
    write_sheet(workbook.create_sheet("Summary"), summary_rows)
    write_sheet(workbook.create_sheet("Evidence"), evidence_rows, widths={1: 34, 5: 110})
    workbook.save(xlsx_out)
    return csv_out, xlsx_out


def print_summary(csv_rows, counts, errors, csv_out, xlsx_out, csv_in):
    print(f"Processed {len(csv_rows)} rows "
          f"({counts['overall'][MATCH]} match, {counts['overall'][MISMATCH]} mismatch, "
          f"{errors} error)")
    print(f"Input: {csv_in}")
    print(f"CSV  : {csv_out}")
    print(f"Excel: {xlsx_out}")


def find_default_csv():
    exact = BASE_DIR / DEFAULT_CSV_NAME
    if exact.exists():
        return exact
    # Browser downloads are saved as "... (1).csv", "... (2).csv"; use the newest copy.
    copies = sorted(BASE_DIR.glob("oracle_host_memory_specification*.csv"), key=lambda p: p.stat().st_mtime)
    return copies[-1] if copies else exact


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--csv", type=Path, default=None)
    parser.add_argument("--pdf-dir", type=Path, default=BASE_DIR / "Oracle pdfs")
    parser.add_argument("--out-dir", type=Path, default=BASE_DIR / "output")
    args = parser.parse_args()
    args.csv = args.csv or find_default_csv()

    if not args.csv.is_file():
        parser.error(f"input CSV not found: {args.csv}")
    if not args.pdf_dir.is_dir():
        parser.error(f"PDF folder not found: {args.pdf_dir}")

    pdf_index = {p.name.lower(): p for p in args.pdf_dir.rglob("*.pdf")}
    with args.csv.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        missing_columns = [c for c in REQUIRED_COLUMNS if c not in (reader.fieldnames or [])]
        if missing_columns:
            parser.error(f"input CSV {args.csv.name} is missing required column(s): {', '.join(missing_columns)}")
        csv_rows = list(reader)

    result_rows, evidence_rows, counts, errors = build_report_rows(csv_rows, pdf_index)
    summary_rows = build_summary_rows(csv_rows, counts, errors)
    csv_out, xlsx_out = write_reports(args.out_dir, result_rows, summary_rows, evidence_rows)
    print_summary(csv_rows, counts, errors, csv_out, xlsx_out, args.csv)


if __name__ == "__main__":
    main()
