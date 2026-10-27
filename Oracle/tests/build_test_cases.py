"""Build tests/Oracle_Memory_Spec_Test_Cases.xlsx, the QA test-case document for compare_memory_specs.py.

Every test case is executed here, so Actual Result / Status are the observed outcome.
Run from the project folder after changing the script or the tests:

    python tests/build_test_cases.py

If LibreOffice is installed, the summary formulas are calculated before saving so totals show in any viewer;
otherwise Excel calculates them when the file is opened.
"""
import csv
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

PROJECT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT))
sys.path.insert(0, str(PROJECT / "tests"))

import compare_memory_specs as cms  # noqa: E402
import test_compare_memory_specs as t  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402

OUT = PROJECT / "tests" / "Oracle_Memory_Spec_Test_Cases.xlsx"
TODAY = date.today().isoformat()
FIXED_ON = "2026-10-27"  # date DEF-01..DEF-05 were fixed and retested
EXECUTED_BY = "Automated run (tests/build_test_cases.py)"
REF = "tests/test_compare_memory_specs.py::"

def find_soffice():
    for candidate in (shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe",
                      r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if candidate and Path(candidate).exists():
            return candidate
    return None


def recalculate(path):
    """Let LibreOffice compute the formulas so cached totals are stored in the file."""
    soffice = find_soffice()
    if soffice is None:
        return "LibreOffice not found - formulas will be calculated when the file is opened in Excel."
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(path)],
                       capture_output=True, timeout=180)
        converted = Path(tmp) / path.name
        if not converted.exists():
            return "LibreOffice conversion failed - formulas will be calculated when opened in Excel."
        shutil.copyfile(converted, path)
    return "Formulas calculated with LibreOffice."


# ----------------------------------------------------------------------------- check helpers

def ok(actual, passed):
    return actual, "Pass" if passed else "Fail"


def status_check(result, expected_status, must_contain=None):
    good = result.status == expected_status and (must_contain is None or must_contain in result.comment)
    text = f"Status = {result.status}; PDF value = '{result.pdf_value}'"
    if result.comment:
        text += f"; Comment: {result.comment}"
    return ok(text, good)


def run_row(name, field_name=None, value=None, drop=None):
    row = t.csv_row(name)
    if field_name:
        row[field_name] = value
    if drop:
        del row[drop]
    return cms.compare_row(row, t.PDF_INDEX[name.lower()])


def run_script(*args):
    return subprocess.run([sys.executable, str(t.SCRIPT), *args], capture_output=True, text=True,
                          timeout=300, cwd=str(PROJECT))


SPECS_CACHE = {}


def specs_for(name):
    if name not in SPECS_CACHE:
        SPECS_CACHE[name] = cms.extract_pdf(t.PDF_INDEX[name.lower()])
    return SPECS_CACHE[name]


def actual_pdf_values(name):
    s = specs_for(name)
    _, total, pp, _, derived = cms.slot_values(s)
    mem, _ = cms.select_memory_value(cms.distinct(s.memory_spec), cms.distinct(s.memory_summary))
    return (total if total is not None else derived, pp, cms.format_size(mem),
            t.largest(s.dimm["LRDIMM"]), t.largest(s.dimm["RDIMM"]),
            t.largest(s.dimm["UDIMM"]), t.largest(s.dimm["SODIMM"]))


# ----------------------------------------------------------------------------- checks

def c_parse_list():
    got = [cms.parse_csv_cell("['24']"), cms.parse_csv_cell("['1.5 TB']"), cms.parse_csv_cell("['16 GB', '32 GB']")]
    return ok(f"{got[0]}, {got[1]}, {got[2]}", got == [["24"], ["1.5 TB"], ["16 GB", "32 GB"]])


def c_missing_markers():
    raws = ["['nan']", "nan", "", None, "['none']", "n/a", "-"]
    got = {repr(r): cms.parse_csv_cell(r) for r in raws}
    return ok("All returned [] (empty): " + ", ".join(got), all(v == [] for v in got.values()))


def c_size():
    got = (cms.parse_csv_size("768 GB"), cms.parse_csv_size("1.5 TB"), cms.parse_csv_size("6 tb"))
    return ok(f"768 GB -> {got[0]:g} GB; 1.5 TB -> {got[1]:g} GB; 6 tb -> {got[2]:g} GB", got == (768, 1536, 6144))


def c_bad_size():
    try:
        cms.parse_csv_size("lots of memory")
    except ValueError as exc:
        return ok(f"ValueError raised: {exc}", "Unrecognised size" in str(exc))
    return ok("No error raised", False)


def c_words():
    got = [cms.word_to_int(w) for w in ("Twenty-four", "Ninety-six", "sixteen", "192")]
    return ok(f"Twenty-four={got[0]}, Ninety-six={got[1]}, sixteen={got[2]}, 192={got[3]}", got == [24, 96, 16, 192])


def c_rdimm_lrdimm():
    s = t.specs_from_text("• RDIMM options: 16 GB at DDR4-2666 and 32 GB at DDR4-2666 • LRDIMM option: 64 GB at DDR4-2666")
    r, lr = cms.distinct(s.dimm["RDIMM"]), cms.distinct(s.dimm["LRDIMM"])
    return ok(f"RDIMM sizes = {[int(v) for v in r]}; LRDIMM sizes = {[int(v) for v in lr]}", r == [16, 32] and lr == [64])


def c_broken_bullet():
    s = t.specs_from_text("\ufffd RDIMM options: 16 GB at DDR4-2666 single rank \ufffd LRDIMM option: 64 GB at "
                          "DDR4-2666 dual rank Oracle Server X8-2 is the most versatile 1.5 TB")
    lr, r = cms.distinct(s.dimm["LRDIMM"]), cms.distinct(s.dimm["RDIMM"])
    return ok(f"LRDIMM sizes = {[int(v) for v in lr]}; RDIMM sizes = {[int(v) for v in r]}", lr == [64] and r == [16])


def c_ddr_speed():
    s = t.specs_from_text("16 dual inline memory module (DIMM) slots per processor support half and fully populated "
                          "memory configurations using 16, 32, or 64 GB DDR4-2400 DIMMs")
    tot, pp = cms.distinct(s.slots_total), cms.distinct(s.slots_per_proc)
    return ok(f"Total slots found = {tot}; per-processor = {[int(v) for v in pp]}", tot == [] and pp == [16])


def c_per_channel():
    s = t.specs_from_text("1,333 MT/sec at three DIMMs per channel (3 DPC)")
    found = s.slots_total + s.slots_per_proc
    return ok(f"Slot values found = {[f.value for f in found]}", found == [])


def c_word_slots():
    s = t.specs_from_text("Ninety-six DIMM slots (12 per CPU module) provide up to 6 TB of DDR4 ECC DIMM memory")
    got = (cms.distinct(s.slots_total), cms.distinct(s.slots_per_proc), cms.distinct(s.memory_spec))
    return ok(f"Total = {[int(v) for v in got[0]]}; per CPU = {[int(v) for v in got[1]]}; "
              f"max memory = {cms.format_size(got[2][0]) if got[2] else '-'}", got == ([96], [12], [6144]))


def c_spec_vs_summary():
    s = t.specs_from_text("• Twenty-four DIMM slots with maximum memory of 1.5 TB "
                          "• Twenty-four DIMM slots provide up to 768 GB of DDR4 DIMM memory")
    sel = cms.select_memory_value(cms.distinct(s.memory_spec), cms.distinct(s.memory_summary))
    return ok(f"Spec = {[cms.format_size(v) for v in cms.distinct(s.memory_spec)]}; summary = "
              f"{[cms.format_size(v) for v in cms.distinct(s.memory_summary)]}; selected = {cms.format_size(sel[0])} ({sel[1]})",
              sel == (768, "spec table"))


def c_procs():
    s = t.specs_from_text("Oracle's SPARC T8-4 server is a resilient, four-processor system")
    got = cms.distinct(s.max_processors)
    return ok(f"Processor count = {[int(v) for v in got]}", got == [4])


def rdimm(*sizes):
    s = cms.PdfSpecs()
    s.dimm["RDIMM"] = t.findings(*sizes)
    return s


def mem(spec=(), summary=()):
    s = cms.PdfSpecs()
    s.memory_spec = t.findings(*spec, kind="spec")
    s.memory_summary = t.findings(*summary, kind="summary")
    return s


def slots(total=(), per_proc=(), processors=()):
    s = cms.PdfSpecs()
    s.slots_total = t.findings(*total, kind="total")
    s.slots_per_proc = t.findings(*per_proc, kind="per_processor")
    s.max_processors = t.findings(*processors)
    return s


def c_overall():
    good = cms.FieldResult("", "", cms.MATCH, "", [])
    bad = cms.FieldResult("", "", cms.MISMATCH, "", [])
    a, b = cms.overall({"a": good, "b": good}), cms.overall({"a": good, "b": bad})
    return ok(f"All fields Match -> {a}; one field Mismatch -> {b}", (a, b) == (cms.MATCH, cms.MISMATCH))


def c_slots_pp():
    r = cms.compare_slots(["16"], slots(per_proc=[16], processors=[4]))
    return ok(f"Status = {r.status}; PDF value = '{r.pdf_value}'", r.status == cms.MATCH and "64 total" in r.pdf_value)


def c_columns():
    with t.CSV_PATH.open(newline="", encoding="utf-8-sig") as fh:
        header = next(csv.reader(fh))
    missing = [c for c in cms.FIELDS + ["file_name"] if c not in header]
    return ok(f"Header = {header}; missing = {missing or 'none'}", not missing)


def c_rows_have_pdf():
    rows = t.read_csv_rows()
    missing = [r["file_name"] for r in rows if r["file_name"].strip().lower() not in t.PDF_INDEX]
    return ok(f"{len(rows)} rows checked; PDFs not found: {missing or 'none'}", not missing)


def c_golden():
    failed = [n for n, exp in t.EXPECTED_PDF_VALUES.items() if actual_pdf_values(n) != exp + ("", "")]
    return ok(f"{len(t.EXPECTED_PDF_VALUES) - len(failed)} of {len(t.EXPECTED_PDF_VALUES)} PDFs match expected values"
              + (f"; failed: {failed}" if failed else " (see 'Test Data' sheet)"), not failed)


def c_all_rows_match():
    rows = t.read_csv_rows()
    bad = []
    for row in rows:
        results, error = cms.compare_row(row, t.PDF_INDEX.get(row["file_name"].lower()))
        if error or cms.overall(results) != cms.MATCH:
            bad.append(row["file_name"])
    return ok(f"{len(rows) - len(bad)} of {len(rows)} rows overall = Match" + (f"; not matching: {bad}" if bad else ""),
              not bad)


def c_wrong(name, field_name, value):
    def check():
        results, error = run_row(name, field_name, value)
        r = results[field_name]
        return ok(f"{field_name} status = {r.status}; overall = {cms.overall(results)}; Comment: {r.comment}",
                  not error and r.status == cms.MISMATCH and cms.overall(results) == cms.MISMATCH)
    return check


def c_missing_pdf():
    results, error = cms.compare_row(t.csv_row("x5-2-datasheet.pdf"), None)
    return ok(f"Result = {results}; message = '{error}'", results is None and "not found" in error)


def c_bad_size_row():
    row = t.csv_row("x5-2-datasheet.pdf")
    row["maximum_memory"] = "['a lot']"
    _, results, error = cms.safe_compare_row(row, t.PDF_INDEX)
    return ok(f"Row marked Error; message = '{error}'", results is None and "Unrecognised size" in error)


def c_missing_column():
    results, error = run_row("x5-2-datasheet.pdf", drop="maximum_UDIMM")
    return ok(f"Row reported as Error; message = '{error}'",
              results is None and "missing column(s): maximum_UDIMM" in error)


def c_missing_column_file():
    with tempfile.TemporaryDirectory() as tmp:
        bad = Path(tmp) / "no_udimm.csv"
        with t.CSV_PATH.open(newline="", encoding="utf-8-sig") as src, bad.open("w", newline="", encoding="utf-8") as dst:
            reader = csv.DictReader(src)
            writer = csv.DictWriter(dst, fieldnames=[c for c in reader.fieldnames if c != "maximum_UDIMM"],
                                    extrasaction="ignore")
            writer.writeheader()
            writer.writerows(reader)
        run = run_script("--csv", str(bad), "--out-dir", tmp)
        made = (Path(tmp) / "memory_spec_comparison.csv").exists()
    last = run.stderr.strip().splitlines()[-1] if run.stderr.strip() else ""
    return ok(f"Exit code {run.returncode}; {last}; report created = {made}",
              run.returncode == 2 and "missing required column(s): maximum_UDIMM" in last and not made)


def c_missing_pdf_dir():
    with tempfile.TemporaryDirectory() as tmp:
        run = run_script("--csv", str(t.CSV_PATH), "--pdf-dir", str(Path(tmp) / "none"), "--out-dir", tmp)
    last = run.stderr.strip().splitlines()[-1] if run.stderr.strip() else ""
    return ok(f"Exit code {run.returncode}; {last}", run.returncode == 2 and "PDF folder not found" in last)


E2E = {}


def e2e():
    if not E2E:
        tmp = tempfile.mkdtemp()
        run = run_script("--csv", str(t.CSV_PATH), "--out-dir", tmp)
        E2E.update(run=run, dir=Path(tmp))
    return E2E["run"], E2E["dir"]


def c_e2e_run():
    run, out = e2e()
    files = sorted(p.name for p in out.iterdir())
    line = run.stdout.strip().splitlines()[0] if run.stdout.strip() else run.stderr.strip()[-200:]
    n = len(t.read_csv_rows())
    return ok(f"Exit code {run.returncode}; console: '{line}'; files: {files}",
              run.returncode == 0 and f"({n} match, 0 mismatch" in run.stdout and len(files) == 2)


def c_e2e_excel():
    _, out = e2e()
    wb = load_workbook(out / "memory_spec_comparison.xlsx")
    res, ev = wb["Results"].max_row, wb["Evidence"].max_row
    names = wb.sheetnames
    summary_header = [c.value for c in wb["Summary"][1]]
    wb.close()
    n = len(t.read_csv_rows())
    return ok(f"Sheets = {names}; Results rows = {res - 1} (+header); Evidence rows = {ev - 1}; "
              f"Summary columns = {summary_header}",
              names == ["Results", "Summary", "Evidence"] and res == n + 1 and ev > n
              and summary_header == ["field", "Match", "Mismatch"])


def run_unittest(test_id):
    import unittest as ut
    suite = ut.defaultTestLoader.loadTestsFromName(f"test_compare_memory_specs.{test_id}")
    result = ut.TestResult()
    suite.run(result)
    problems = result.failures + result.errors
    return (result.testsRun, problems)


def c_readme_lines():
    import re
    readme = (PROJECT / "README.md").read_text(encoding="utf-8")
    rows = t.TestReadme.ROW.findall(readme)
    runs, problems = run_unittest("TestReadme.test_walkthrough_line_numbers_match_the_code")
    ex = [r for r in rows if r[2] == "extract_pdf"]
    return ok(f"{len(rows)} walkthrough rows checked; failures = {len(problems)}; "
              f"e.g. extract_pdf documented at lines {ex[0][0]}-{ex[0][1]}", runs == 1 and not problems)


def c_readme_complete():
    runs, problems = run_unittest("TestReadme.test_every_top_level_definition_is_documented")
    return ok("Every top-level function, class and constant has a walkthrough row" if not problems
              else f"Undocumented: {problems[0][1].splitlines()[-1]}", runs == 1 and not problems)


def c_requirements():
    from importlib.metadata import version
    pins = (PROJECT / "requirements.txt").read_text(encoding="utf-8").split()
    installed = [f"{p.split('==')[0]}=={version(p.split('==')[0])}" for p in pins]
    return ok(f"requirements.txt: {pins}; installed: {installed}", pins == installed and len(pins) == 2)


def c_missing_csv():
    with tempfile.TemporaryDirectory() as tmp:
        run = run_script("--csv", str(Path(tmp) / "nope.csv"), "--out-dir", tmp)
    last = run.stderr.strip().splitlines()[-1] if run.stderr.strip() else ""
    return ok(f"Exit code {run.returncode}; {last}",
              run.returncode == 2 and "input CSV not found" in last and "Traceback" not in run.stderr)


def c_default_run():
    with tempfile.TemporaryDirectory() as tmp:
        run = run_script("--out-dir", tmp)
    out = run.stdout.strip().splitlines()
    shown = " | ".join(out[:2]) if run.returncode == 0 else run.stderr.strip().splitlines()[-1]
    return ok(f"Exit code {run.returncode}; {shown}", run.returncode == 0 and f"Input: {t.CSV_PATH}" in run.stdout)


# ----------------------------------------------------------------------------- test-case catalogue

PRE_UNIT = "compare_memory_specs.py importable; Python 3.9+ with pymupdf and openpyxl installed"
PRE_DATA = "Input CSV and the 26 datasheet PDFs present in C:\\Automation\\oracle"
CASES = []


def tc(module, req, scenario, title, ttype, prio, pre, data, steps, expected, check, auto, remarks=""):
    CASES.append(dict(module=module, req=req, scenario=scenario, title=title, type=ttype, prio=prio, pre=pre,
                      data=data, steps=steps, expected=expected, check=check, auto=auto, remarks=remarks))


M1, M2, M3, M4, M5, M6, M7 = ("CSV Input Parsing", "PDF Extraction", "Comparison Rules", "Data Validation",
                              "Negative / Error Handling", "Reporting (End-to-End)", "Known Issues")
M8 = "Documentation & Setup"

tc(M1, "REQ-07", "Read list-formatted CSV cells", "Verify list-style cell values are parsed correctly", "Functional",
   "High", PRE_UNIT, "['24'] ; ['1.5 TB'] ; ['16 GB', '32 GB']",
   "1. Call parse_csv_cell() with each test value\n2. Observe returned list",
   "Returns ['24'], ['1.5 TB'] and ['16 GB', '32 GB']", c_parse_list, "TestHelpers.test_parse_csv_cell")
tc(M1, "REQ-04", "Missing value markers", "Verify missing-value markers are treated as empty", "Boundary", "High",
   PRE_UNIT, "['nan'] ; nan ; (empty) ; None ; ['none'] ; n/a ; -",
   "1. Call parse_csv_cell() with each marker\n2. Observe returned list",
   "Every marker returns an empty list", c_missing_markers,
   "TestHelpers.test_parse_csv_cell_treats_missing_markers_as_empty")
tc(M1, "REQ-07", "Memory size conversion", "Verify GB/TB values convert to GB (1 TB = 1024 GB)", "Functional",
   "High", PRE_UNIT, "768 GB ; 1.5 TB ; 6 tb", "1. Call parse_csv_size() with each value\n2. Observe GB result",
   "768 / 1536 / 6144 GB (unit is case-insensitive)", c_size, "TestHelpers.test_parse_csv_size")
tc(M1, "REQ-05", "Invalid memory size text", "Verify non-numeric size text is rejected", "Negative", "Medium",
   PRE_UNIT, "lots of memory", "1. Call parse_csv_size('lots of memory')",
   "ValueError 'Unrecognised size value in CSV' is raised", c_bad_size,
   "TestHelpers.test_parse_csv_size_rejects_bad_text")
tc(M1, "REQ-08", "Number words", "Verify spelled-out numbers are converted", "Functional", "Medium", PRE_UNIT,
   "Twenty-four ; Ninety-six ; sixteen ; 192", "1. Call word_to_int() with each value",
   "24, 96, 16, 192", c_words, "TestHelpers.test_word_to_int")

tc(M2, "REQ-03", "RDIMM vs LRDIMM text", "Verify RDIMM pattern does not match inside 'LRDIMM'", "Functional",
   "High", PRE_UNIT, "• RDIMM options: 16 GB ... and 32 GB ... • LRDIMM option: 64 GB ...",
   "1. Extract DIMM values from the sample text\n2. Compare RDIMM and LRDIMM lists",
   "RDIMM = [16, 32]; LRDIMM = [64]", c_rdimm_lrdimm, "TestLocators.test_rdimm_is_not_confused_with_lrdimm")
tc(M2, "REQ-08", "Broken bullet glyph (X8-2 datasheet)",
   "Verify LRDIMM is found when bullets are U+FFFD and the line runs into sidebar text", "Regression", "High",
   PRE_UNIT, "\\ufffd RDIMM options: 16 GB ... \\ufffd LRDIMM option: 64 GB ... Oracle Server X8-2 ... 1.5 TB",
   "1. Extract DIMM values from the sample text", "LRDIMM = [64]; RDIMM = [16]; 1.5 TB not picked up",
   c_broken_bullet, "TestLocators.test_broken_bullet_and_run_on_sidebar_text",
   "Regression for bug found during development (X8-2 LRDIMM missing)")
tc(M2, "REQ-01", "DDR speed next to 'DIMMs'", "Verify 'DDR4-2400 DIMMs' is not read as 2400 slots", "Regression",
   "High", PRE_UNIT, "... using 16, 32, or 64 GB DDR4-2400 DIMMs",
   "1. Extract slot values from the sample text", "No total slots; per-processor slots = [16]", c_ddr_speed,
   "TestLocators.test_ddr_speed_is_not_read_as_slot_count",
   "Regression for bug found during development (M12-1 showed 2400 slots)")
tc(M2, "REQ-01", "DIMMs per channel", "Verify 'three DIMMs per channel' is not a slot count", "Negative", "Medium",
   PRE_UNIT, "1,333 MT/sec at three DIMMs per channel (3 DPC)", "1. Extract slot values from the sample text",
   "No slot values found", c_per_channel, "TestLocators.test_dimms_per_channel_is_not_a_slot_count")
tc(M2, "REQ-01", "Word numbers and per-CPU slots", "Verify word-number totals and '(12 per CPU module)'",
   "Functional", "High", PRE_UNIT, "Ninety-six DIMM slots (12 per CPU module) provide up to 6 TB ...",
   "1. Extract slots and memory from the sample text", "Total = 96; per CPU = 12; max memory = 6 TB",
   c_word_slots, "TestLocators.test_word_numbers_and_per_cpu_module")
tc(M2, "REQ-02", "Spec table vs summary text", "Verify spec-table and summary memory values are kept apart",
   "Functional", "High", PRE_UNIT,
   "• ... slots with maximum memory of 1.5 TB • ... DIMM slots provide up to 768 GB ...",
   "1. Extract memory values\n2. Call select_memory_value()", "Selected = 768 GB from 'spec table'",
   c_spec_vs_summary, "TestLocators.test_spec_table_and_summary_are_kept_apart")
tc(M2, "REQ-01", "Processor count", "Verify processor count is read from 'four-processor system'", "Functional",
   "Medium", PRE_UNIT, "Oracle's SPARC T8-4 server is a resilient, four-processor system",
   "1. Extract processor counts from the sample text", "Processor count = [4]", c_procs,
   "TestLocators.test_processor_count")

tc(M3, "REQ-04", "DIMM empty on both sides", "Verify DIMM field empty in CSV and PDF is Match", "Boundary", "Medium",
   PRE_UNIT, "CSV RDIMM = (empty); PDF RDIMM = none", "1. Call compare_dimm('RDIMM', [], empty PDF specs)",
   "Status = Match", lambda: status_check(cms.compare_dimm("RDIMM", [], cms.PdfSpecs()), cms.MATCH),
   "TestComparisonRules.test_dimm_both_empty_is_match")
tc(M3, "REQ-03", "Largest DIMM option", "Verify CSV is compared with the largest PDF option", "Functional", "High",
   PRE_UNIT, "CSV RDIMM = 32 GB; PDF RDIMM options = 16 GB, 32 GB", "1. Call compare_dimm()",
   "Status = Match; PDF value = 32 GB",
   lambda: status_check(cms.compare_dimm("RDIMM", ["32 GB"], rdimm(16, 32)), cms.MATCH),
   "TestComparisonRules.test_dimm_largest_pdf_option_is_compared")
tc(M3, "REQ-03", "Smaller DIMM option in CSV", "Verify a non-maximum DIMM size is a Mismatch", "Negative", "High",
   PRE_UNIT, "CSV RDIMM = 16 GB; PDF RDIMM options = 16 GB, 32 GB", "1. Call compare_dimm()",
   "Status = Mismatch; comment lists 'PDF options: 16 GB, 32 GB'",
   lambda: status_check(cms.compare_dimm("RDIMM", ["16 GB"], rdimm(16, 32)), cms.MISMATCH,
                        "PDF options: 16 GB, 32 GB"),
   "TestComparisonRules.test_dimm_different_value_is_mismatch")
tc(M3, "REQ-04", "DIMM missing in CSV", "Verify DIMM present only in PDF is a Mismatch", "Negative", "High",
   PRE_UNIT, "CSV RDIMM = (empty); PDF RDIMM = 32 GB", "1. Call compare_dimm()",
   "Status = Mismatch; comment says 'missing in CSV'",
   lambda: status_check(cms.compare_dimm("RDIMM", [], rdimm(32)), cms.MISMATCH, "missing in CSV"),
   "TestComparisonRules.test_dimm_missing_on_one_side_is_mismatch")
tc(M3, "REQ-04", "DIMM missing in PDF", "Verify DIMM present only in CSV is a Mismatch", "Negative", "High",
   PRE_UNIT, "CSV RDIMM = 32 GB; PDF RDIMM = none", "1. Call compare_dimm()",
   "Status = Mismatch; comment says PDF does not mention RDIMM",
   lambda: status_check(cms.compare_dimm("RDIMM", ["32 GB"], cms.PdfSpecs()), cms.MISMATCH, "does not mention"),
   "TestComparisonRules.test_dimm_missing_on_one_side_is_mismatch")
tc(M3, "REQ-02", "Spec table wins over summary", "Verify CSV equal to spec table is Match even if summary differs",
   "Functional", "High", PRE_UNIT, "CSV = 768 GB; PDF spec = 768 GB; PDF summary = 1.5 TB",
   "1. Call compare_memory()", "Status = Match; comment notes PDF is 'internally inconsistent'",
   lambda: status_check(cms.compare_memory(["768 GB"], mem([768], [1536])), cms.MATCH, "internally inconsistent"),
   "TestComparisonRules.test_memory_spec_table_wins_and_conflict_is_noted", "Business rule chosen by the user on 2026-10-26 (spec table wins)")
tc(M3, "REQ-02", "CSV equals summary only", "Verify CSV equal only to summary value is Mismatch", "Negative",
   "High", PRE_UNIT, "CSV = 1.5 TB; PDF spec = 768 GB; PDF summary = 1.5 TB", "1. Call compare_memory()",
   "Status = Mismatch; comment says value 'appears elsewhere' in the PDF",
   lambda: status_check(cms.compare_memory(["1.5 TB"], mem([768], [1536])), cms.MISMATCH, "appears elsewhere"),
   "TestComparisonRules.test_memory_matching_only_the_summary_is_mismatch")
tc(M3, "REQ-02", "Several memory configurations", "Verify the largest spec-table configuration is used",
   "Functional", "Medium", PRE_UNIT, "CSV = 8 TB; PDF spec = 4 TB and 8 TB", "1. Call compare_memory()",
   "Status = Match; comment mentions 'smaller configuration'",
   lambda: status_check(cms.compare_memory(["8 TB"], mem([4096, 8192])), cms.MATCH, "smaller configuration"),
   "TestComparisonRules.test_memory_largest_configuration_is_used")
tc(M3, "REQ-04", "Memory missing in CSV", "Verify empty CSV memory with PDF value is Mismatch", "Negative", "High",
   PRE_UNIT, "CSV = (empty); PDF spec = 16 TB", "1. Call compare_memory()",
   "Status = Mismatch; comment says 'missing in CSV'",
   lambda: status_check(cms.compare_memory([], mem([16384])), cms.MISMATCH, "missing in CSV"),
   "TestComparisonRules.test_memory_missing_in_csv_is_mismatch")
tc(M3, "REQ-01", "Per-processor slot count", "Verify CSV per-processor count is Match and total is shown",
   "Functional", "High", PRE_UNIT, "CSV = 16; PDF = 16 per processor, 4 processors", "1. Call compare_slots()",
   "Status = Match; PDF value shows '64 total'", c_slots_pp,
   "TestComparisonRules.test_slots_per_processor_and_derived_total_both_match")
tc(M3, "REQ-01", "Derived total slot count", "Verify CSV equal to derived total (per-proc x processors) is Match",
   "Functional", "High", PRE_UNIT, "CSV = 64; PDF = 16 per processor, 4 processors", "1. Call compare_slots()",
   "Status = Match",
   lambda: status_check(cms.compare_slots(["64"], slots(per_proc=[16], processors=[4])), cms.MATCH),
   "TestComparisonRules.test_slots_per_processor_and_derived_total_both_match")
tc(M3, "REQ-01", "Wrong slot count", "Verify a slot count that is neither per-processor nor total is Mismatch",
   "Negative", "High", PRE_UNIT, "CSV = 32; PDF = 16 per processor, 4 processors", "1. Call compare_slots()",
   "Status = Mismatch",
   lambda: status_check(cms.compare_slots(["32"], slots(per_proc=[16], processors=[4])), cms.MISMATCH),
   "TestComparisonRules.test_slots_other_number_is_mismatch")
tc(M3, "REQ-06", "Overall row status", "Verify any field Mismatch makes the row Mismatch", "Functional", "High",
   PRE_UNIT, "Fields: (Match, Match) and (Match, Mismatch)", "1. Call overall() for both sets",
   "Match, then Mismatch", c_overall, "TestComparisonRules.test_overall_status")

tc(M4, "REQ-07", "CSV structure", "Verify the input CSV has all required columns", "Data Validation", "High",
   PRE_DATA, "Input CSV header", "1. Open the input CSV\n2. Read the header row",
   "dimm_slots, maximum_memory, maximum_LRDIMM, maximum_RDIMM, maximum_UDIMM, maximum_SODIMM, file_name all present",
   c_columns, "TestCsvAgainstPdfs.test_csv_has_all_required_columns")
tc(M4, "REQ-05", "CSV-to-PDF mapping", "Verify every CSV row has a matching PDF file", "Data Validation", "High",
   PRE_DATA, "file_name column (26 rows)", "1. For each row, look up file_name in the PDF folder",
   "All 26 PDFs found", c_rows_have_pdf, "TestCsvAgainstPdfs.test_every_csv_row_has_a_pdf")
tc(M4, "REQ-08", "PDF extraction accuracy", "Verify extracted values equal hand-verified values for all 26 PDFs",
   "Regression", "High", PRE_DATA, "See 'Test Data' sheet (26 PDFs x 7 fields)",
   "1. Extract every PDF\n2. Compare each field with the hand-verified value",
   "All 26 PDFs match the hand-verified values", c_golden,
   "TestRealPdfs.test_extracted_values_match_hand_verified_values", "Golden data verified manually against PDFs")
tc(M4, "REQ-06", "Full data comparison", "Verify every CSV row matches its PDF", "Data Validation", "High",
   PRE_DATA, "Input CSV (26 rows)", "1. Run compare_row() for every CSV row\n2. Check overall status",
   "26 of 26 rows overall = Match", c_all_rows_match, "TestCsvAgainstPdfs.test_every_csv_row_matches_its_pdf")

for name, fld, val, why in [
    ("x8-2-datasheet.pdf", "maximum_RDIMM", "['16 GB']", "smaller RDIMM option instead of the maximum"),
    ("x8-2-datasheet.pdf", "maximum_LRDIMM", "['32 GB']", "wrong LRDIMM size"),
    ("x5-2-datasheet.pdf", "dimm_slots", "['12']", "wrong slot count"),
    ("sparc-t8-4-ds-3864254.pdf", "dimm_slots", "['32']", "neither per-processor (16) nor total (64)"),
    ("x6-2-datasheet-v4.pdf", "maximum_memory", "['1.5 TB']", "summary value instead of spec-table value"),
    ("sparc-m8-8-ds-3864183.pdf", "maximum_memory", "['nan']", "value missing in CSV"),
    ("x7-2_datasheet-v1.pdf", "maximum_RDIMM", "['32 GB']", "PDF never mentions RDIMM"),
    ("x9-2-datasheet.pdf", "maximum_UDIMM", "['16 GB']", "PDF never mentions UDIMM"),
]:
    tc(M5, {"dimm_slots": "REQ-01", "maximum_memory": "REQ-02"}.get(fld, "REQ-03"),
       f"Wrong {fld} detected", f"Verify {fld} = {val} for {name} is reported ({why})", "Negative", "High",
       PRE_DATA, f"Row {name}; {fld} changed to {val}",
       f"1. Load the {name} row from the CSV\n2. Change {fld} to {val}\n3. Run compare_row()",
       f"{fld} status = Mismatch; overall = Mismatch", c_wrong(name, fld, val),
       "TestCatchesWrongData.test_wrong_values_are_reported_as_mismatch")
tc(M5, "REQ-05", "PDF file missing", "Verify a row whose PDF is missing is reported as Error", "Negative", "High",
   PRE_DATA, "x5-2 row with no PDF path", "1. Call compare_row(row, None)",
   "No results; message 'PDF file not found in the PDF folder.'", c_missing_pdf,
   "TestCatchesWrongData.test_missing_pdf_is_an_error")
tc(M5, "REQ-05", "Invalid data in one row", "Verify bad size text gives an Error row without stopping the run",
   "Negative", "High", PRE_DATA, "x5-2 row; maximum_memory = ['a lot']",
   "1. Change maximum_memory to ['a lot']\n2. Call safe_compare_row()",
   "Row reported as Error with 'Unrecognised size' message; no crash", c_bad_size_row,
   "TestCatchesWrongData.test_bad_size_text_becomes_error_row_not_crash")

tc(M6, "REQ-06", "Script run", "Verify the script runs and writes CSV and Excel reports", "End-to-End", "High",
   PRE_DATA, "--csv <input CSV> --out-dir <temporary folder>",
   "1. Run: python compare_memory_specs.py --csv <input> --out-dir <temp>\n2. Check exit code and console\n"
   "3. Check output folder", "Exit code 0; console shows '26 match, 0 mismatch'; 2 report files created",
   c_e2e_run, "TestEndToEnd.test_script_produces_reports")
tc(M6, "REQ-06", "Excel report structure", "Verify Excel report has Results, Summary and Evidence sheets",
   "End-to-End", "High", PRE_DATA, "memory_spec_comparison.xlsx from the run above",
   "1. Open the Excel report\n2. Check sheet names, row counts and Summary columns",
   "Sheets Results, Summary, Evidence; Results has 26 data rows; Evidence has supporting rows; "
   "Summary columns are field / Match / Mismatch (no unused Review column)", c_e2e_excel,
   "TestEndToEnd.test_script_produces_reports", "Retest for DEF-04 (fixed)")
tc(M8, "REQ-09", "README walkthrough accuracy", "Verify every README walkthrough row points at the right code line",
   "Regression", "Low", "README.md and compare_memory_specs.py present", "README 'Code walkthrough' tables",
   "1. Read each walkthrough row (lines + name)\n2. Check the code at the first line defines that name",
   "Every row's first line is the definition of the named function/constant", c_readme_lines,
   "TestReadme.test_walkthrough_line_numbers_match_the_code", "Retest for DEF-03 (fixed)")
tc(M8, "REQ-09", "README walkthrough completeness", "Verify every top-level definition is documented in the README",
   "Regression", "Low", "README.md and compare_memory_specs.py present", "All top-level definitions in the script",
   "1. List top-level functions, classes and constants\n2. Check each has a walkthrough row",
   "No undocumented definitions", c_readme_complete,
   "TestReadme.test_every_top_level_definition_is_documented", "Retest for DEF-03 (fixed)")
tc(M8, "REQ-09", "Pinned dependencies", "Verify requirements.txt pins the installed library versions",
   "Configuration", "Low", "requirements.txt present", "pymupdf, openpyxl",
   "1. Read requirements.txt\n2. Compare with installed package versions",
   "pymupdf==1.28.2 and openpyxl==3.1.5, equal to the installed versions", c_requirements,
   "TestDependencies.test_requirements_match_installed_versions", "Retest for DEF-05 (fixed)")
tc(M6, "REQ-05", "Input CSV not found", "Verify a missing CSV path gives a clear message, not a traceback",
   "Negative", "Medium", "Script available", "--csv nope.csv", "1. Run the script with a non-existent CSV path",
   "Exit code 2; message 'input CSV not found'; no Python traceback", c_missing_csv,
   "TestEndToEnd.test_missing_csv_file_fails_with_message", "Retest for DEF-02 (fixed)")
tc(M6, "REQ-05", "PDF folder not found", "Verify a missing PDF folder gives a clear message", "Negative", "Medium",
   "Script available", "--pdf-dir <non-existent folder>", "1. Run the script with a non-existent PDF folder",
   "Exit code 2; message 'PDF folder not found'", c_missing_pdf_dir,
   "TestEndToEnd.test_missing_pdf_folder_fails_with_message")
tc(M6, "REQ-06", "Default run (README usage)", "Verify 'python compare_memory_specs.py' works with no options",
   "End-to-End", "Medium", PRE_DATA, "No --csv option; CSV on disk is 'oracle_host_memory_specification (1).csv'",
   "1. Run: python compare_memory_specs.py --out-dir <temp>\n2. Check exit code and the 'Input:' line",
   "Exit code 0; console shows the CSV file that was used; reports created", c_default_run,
   "TestEndToEnd.test_default_run_without_options", "Retest for DEF-02 (fixed)")

tc(M5, "REQ-07", "Missing CSV column (row)", "Verify a row missing a column is reported as Error, not silently matched",
   "Negative", "High", PRE_DATA, "x5-2 row with the maximum_UDIMM column removed",
   "1. Remove maximum_UDIMM from the row\n2. Run compare_row()",
   "Row reported as Error: 'CSV row is missing column(s): maximum_UDIMM'", c_missing_column,
   "TestCatchesWrongData.test_missing_csv_column_is_an_error_not_a_match", "Retest for DEF-01 (fixed)")
tc(M5, "REQ-07", "Missing CSV column (file)", "Verify a CSV file missing a column is rejected before comparing",
   "Negative", "High", PRE_DATA, "Copy of the input CSV without the maximum_UDIMM column",
   "1. Save the CSV without maximum_UDIMM\n2. Run the script with --csv <that file>",
   "Exit code 2; message names maximum_UDIMM; no report created", c_missing_column_file,
   "TestEndToEnd.test_csv_file_missing_a_column_fails_with_message", "Retest for DEF-01 (fixed)")

# ----------------------------------------------------------------------------- execute

for i, case in enumerate(CASES, start=1):
    case["id"] = f"TC_OMS_{i:03d}"
    try:
        case["actual"], case["status"] = case["check"]()
    except Exception as exc:  # a crashing check is a failed test, not a crashed report
        case["actual"], case["status"] = f"Exception: {exc!r}", "Fail"
    print(case["id"], case["status"], "-", case["title"])

# ----------------------------------------------------------------------------- workbook styling

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color="1F4E78")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
FILL = {"Pass": "C6EFCE", "Fail": "FFC7CE", "Blocked": "FFEB9C", "Not Run": "D9D9D9", "Open": "FFC7CE",
        "Covered": "C6EFCE"}


def table(ws, header, rows, widths, start_row=1):
    for c, h in enumerate(header, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r, row in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font, cell.alignment, cell.border = BODY_FONT, WRAP, BORDER
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + c) if c <= 26 else "A" + chr(64 + c - 26)].width = w
    end_col = chr(64 + len(header))
    ws.auto_filter.ref = f"A{start_row}:{end_col}{start_row + len(rows)}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=2)


def status_colours(ws, col_range):
    for value, colour in FILL.items():
        ws.conditional_formatting.add(col_range, CellIsRule(operator="equal", formula=[f'"{value}"'],
                                                            fill=PatternFill("solid", fgColor=colour)))


wb = Workbook()

# --- Test Cases
ws = wb.active
ws.title = "Test Cases"
header = ["TC ID", "Module", "Req ID", "Test Scenario", "Test Case Title", "Test Type", "Priority",
          "Preconditions", "Test Data", "Test Steps", "Expected Result", "Actual Result", "Status",
          "Automation Status", "Automated Test Reference", "Executed By", "Execution Date", "Remarks / Defect ID"]
rows = [[c["id"], c["module"], c["req"], c["scenario"], c["title"], c["type"], c["prio"], c["pre"], c["data"],
         c["steps"], c["expected"], c["actual"], c["status"],
         "Automated" if not c["auto"].startswith("Manual") else "Manual",
         (REF + c["auto"]) if not c["auto"].startswith("Manual") else c["auto"],
         EXECUTED_BY, TODAY, c["remarks"]] for c in CASES]
table(ws, header, rows, [13, 20, 9, 26, 40, 14, 9, 30, 34, 38, 38, 48, 9, 12, 44, 16, 12, 22])
n_cases = len(rows)
status_colours(ws, f"M2:M{n_cases + 1}")
dv = DataValidation(type="list", formula1='"Pass,Fail,Blocked,Not Run"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"M2:M{n_cases + 1}")
pv = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
ws.add_data_validation(pv)
pv.add(f"G2:G{n_cases + 1}")
for r in range(2, n_cases + 2):
    ws.cell(row=r, column=1).font = BOLD
    ws.cell(row=r, column=13).font = BOLD
    ws.cell(row=r, column=13).alignment = Alignment(horizontal="center", vertical="top")

# --- Test Data (golden values)
wd = wb.create_sheet("Test Data")
fields = ["DIMM slots (system)", "Slots per processor", "Max memory", "Max LRDIMM", "Max RDIMM", "Max UDIMM",
          "Max SODIMM"]
dheader = ["#", "PDF file", "CSV row values (input)"]
for f in fields:
    dheader += [f"Expected {f}", f"Actual {f}"]
dheader += ["Result"]
csv_rows = {r["file_name"]: r for r in t.read_csv_rows()}
drows = []
for i, (name, exp) in enumerate(t.EXPECTED_PDF_VALUES.items(), start=1):
    exp = exp + ("", "")
    act = actual_pdf_values(name)
    r = csv_rows.get(name, {})
    csv_text = "; ".join(f"{k}={', '.join(cms.parse_csv_cell(r.get(k))) or '-'}" for k in cms.FIELDS)
    row = [i, name, csv_text]
    for e, a in zip(exp, act):
        row += ["-" if e in ("", None) else e, "-" if a in ("", None) else a]
    drows.append(row)
table(wd, dheader, drows, [5, 32, 46] + [12] * 14 + [10])
last = len(drows) + 1
for r in range(2, last + 1):
    pairs = "*".join(f"({chr(64 + c)}{r}={chr(64 + c + 1)}{r})" for c in range(4, 18, 2))
    wd.cell(row=r, column=18, value=f'=IF({pairs},"Pass","Fail")').font = BOLD
    wd.cell(row=r, column=18).border = BORDER
status_colours(wd, f"R2:R{last}")
wd.cell(row=last + 2, column=2, value="Expected values were read manually from each PDF; Actual values are "
                                      "what compare_memory_specs.py extracted on the execution date. "
                                      "'-' = not stated in the PDF.").font = Font(name=FONT, size=9, italic=True)

# --- Traceability matrix
wt = wb.create_sheet("Traceability Matrix")
reqs = [
    ("REQ-01", "DIMM slot count matches total system slots or per-processor slots"),
    ("REQ-02", "Maximum memory compared with spec-table value (spec table wins over summary text)"),
    ("REQ-03", "LRDIMM / RDIMM / UDIMM / SODIMM compared with the largest size listed in the PDF"),
    ("REQ-04", "Missing values: empty on both sides = Match; present on one side only = Mismatch"),
    ("REQ-05", "Errors (missing PDF, unreadable data, bad input) are reported without stopping the run"),
    ("REQ-06", "Results written to CSV and Excel (Results, Summary, Evidence) with overall status"),
    ("REQ-07", "Input CSV format (list-style cells, required columns) is read correctly"),
    ("REQ-08", "PDF text extraction is accurate across datasheet wording variations"),
    ("REQ-09", "Project is maintainable: README matches the code and dependencies are pinned"),
]
trows = []
for req, desc in reqs:
    ids = ", ".join(c["id"] for c in CASES if c["req"] == req)
    trows.append([req, desc, ids, None, None, None, None])
table(wt, ["Req ID", "Requirement", "Linked Test Cases", "No. of Test Cases", "Passed", "Failed", "Coverage"],
      trows, [9, 60, 48, 12, 10, 10, 12])
for r in range(2, len(trows) + 2):
    tcr = f"'Test Cases'!$C$2:$C${n_cases + 1}"
    tcs = f"'Test Cases'!$M$2:$M${n_cases + 1}"
    wt.cell(row=r, column=4, value=f"=COUNTIF({tcr},A{r})")
    wt.cell(row=r, column=5, value=f'=COUNTIFS({tcr},A{r},{tcs},"Pass")')
    wt.cell(row=r, column=6, value=f'=COUNTIFS({tcr},A{r},{tcs},"Fail")')
    wt.cell(row=r, column=7, value=f'=IF(D{r}>0,"Covered","Not Covered")')
    for c in range(4, 8):
        wt.cell(row=r, column=c).font, wt.cell(row=r, column=c).border = BODY_FONT, BORDER
        wt.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="top")
status_colours(wt, f"G2:G{len(trows) + 1}")
wt.conditional_formatting.add(f"G2:G{len(trows) + 1}", CellIsRule(operator="equal", formula=['"Not Covered"'],
                              fill=PatternFill("solid", fgColor="FFC7CE")))

# --- Defects
wdf = wb.create_sheet("Defects")
def linked(defect_id):
    return ", ".join(c["id"] for c in CASES if defect_id in c["remarks"]) or "-"


defects = [
    ("DEF-01", "Missing CSV column is silently treated as empty and reported as Match", "Medium", "Medium",
     linked("DEF-01"), "1. Remove the maximum_UDIMM column from the input CSV\n2. Run the script",
     "Column reported as missing (Error)", "All rows showed maximum_UDIMM = Match",
     "compare_row() used row.get(key); a missing key became an empty value",
     "Validate the CSV header against the required columns before comparing", "Fixed",
     f"{FIXED_ON}: header validated in main() (exit 2 with column names) and per row in compare_row(). "
     "Retest passed."),
    ("DEF-02", "Default run fails: default CSV file name does not exist", "Low", "Medium",
     linked("DEF-02"), "1. Run: python compare_memory_specs.py (no options)",
     "Report generated, as documented in README", "FileNotFoundError traceback; no report",
     "Default was oracle_host_memory_specification.csv; file on disk is "
     "'oracle_host_memory_specification (1).csv'",
     "Fall back to the downloaded copy; show a friendly error instead of a traceback", "Fixed",
     f"{FIXED_ON}: find_default_csv() uses the newest '(N)' copy and prints "
     "'Input:'; missing CSV / PDF folder give exit 2 with a message. Retest passed."),
    ("DEF-03", "README line-by-line walkthrough references outdated line numbers", "Low", "Low", linked("DEF-03"),
     "1. Open README.md 'Code walkthrough'\n2. Compare with compare_memory_specs.py",
     "Line numbers match the code", "e.g. README said extract_pdf starts at line 197; it was at line 260",
     "Walkthrough was written before the script was refactored", "Regenerate the walkthrough after code changes",
     "Fixed", f"{FIXED_ON}: walkthrough rewritten per definition (61 rows, current line ranges); "
     "TestReadme now fails if any row drifts or a definition is undocumented. Retest passed."),
    ("DEF-04", "'Review' status can no longer occur but is still shown in reports", "Low", "Low", linked("DEF-04"),
     "1. Run the script\n2. Open the Summary sheet", "Only reachable statuses are shown",
     "Review column was always 0", "No compare function returned REVIEW",
     "Remove Review from the summary, or reintroduce a rule that uses it", "Fixed",
     f"{FIXED_ON}: REVIEW removed from statuses, Excel colours, Summary sheet and console output "
     "('26 match, 0 mismatch, 0 error'). Retest passed."),
    ("DEF-05", "Python dependencies are not pinned (no requirements.txt)", "Low", "Low", linked("DEF-05"),
     "1. Look for requirements.txt in the project", "Dependency versions recorded",
     "No requirements file; README listed packages without versions", "Project folder listing",
     "Add requirements.txt (pymupdf==1.28.2, openpyxl==3.1.5)", "Fixed",
     f"{FIXED_ON}: requirements.txt added; README install step uses 'pip install -r requirements.txt'. Retest passed."),
]
table(wdf, ["Defect ID", "Summary", "Severity", "Priority", "Linked TC", "Steps to Reproduce", "Expected",
            "Actual (when found)", "Evidence", "Suggested Fix", "Status", "Resolution / Retest"],
      [list(d) for d in defects], [10, 42, 10, 10, 15, 36, 30, 34, 44, 38, 9, 50])
status_colours(wdf, f"K2:K{len(defects) + 1}")
wdf.conditional_formatting.add(f"K2:K{len(defects) + 1}", CellIsRule(operator="equal", formula=['"Fixed"'],
                               fill=PatternFill("solid", fgColor="C6EFCE")))

# --- Test Summary (first sheet)
wsum = wb.create_sheet("Test Summary", 0)
wsum.column_dimensions["A"].width = 30
wsum.column_dimensions["B"].width = 62
wsum.column_dimensions["C"].width = 14
wsum.column_dimensions["D"].width = 14
wsum["A1"] = "Test Case Document - Oracle Memory Specification Comparator"
wsum["A1"].font = TITLE
info = [
    ("Project", "Oracle host memory specification validation"),
    ("Application under test", "compare_memory_specs.py (CSV vs datasheet PDF comparison)"),
    ("Test objective", "Verify DIMM slots, maximum memory and LRDIMM/RDIMM/UDIMM/SODIMM values in the input "
                       "CSV are correctly compared with the Oracle/Fujitsu datasheet PDFs"),
    ("Test environment", f"Windows 11 Pro; Python {sys.version.split()[0]}; PyMuPDF {__import__('pymupdf').__version__}; "
                         f"openpyxl {__import__('openpyxl').__version__}"),
    ("Test data", f"{t.CSV_PATH.name} (26 rows); 26 datasheet PDFs in 'Oracle pdfs'"),
    ("Automation suite", "tests/test_compare_memory_specs.py  (run: python -m unittest discover -s tests -v)"),
    ("Execution date", TODAY),
    ("Executed by", EXECUTED_BY),
    ("Prepared by", ""),
    ("Reviewed by", ""),
]
for i, (k, v) in enumerate(info, start=3):
    wsum.cell(row=i, column=1, value=k).font = BOLD
    wsum.cell(row=i, column=2, value=v).font = BODY_FONT
    wsum.cell(row=i, column=2).alignment = WRAP
r0 = 3 + len(info) + 1
wsum.cell(row=r0, column=1, value="Execution Summary").font = Font(name=FONT, size=12, bold=True, color="1F4E78")
st = f"'Test Cases'!$M$2:$M${n_cases + 1}"
summary = [
    ("Total test cases", f"=COUNTA('Test Cases'!$A$2:$A${n_cases + 1})"),
    ("Passed", f'=COUNTIF({st},"Pass")'),
    ("Failed", f'=COUNTIF({st},"Fail")'),
    ("Blocked", f'=COUNTIF({st},"Blocked")'),
    ("Not run", f'=COUNTIF({st},"Not Run")'),
    ("Pass rate", f"=IF(B{r0 + 1}>0,B{r0 + 2}/B{r0 + 1},0)"),
    ("Automated test cases", f"=COUNTIF('Test Cases'!$N$2:$N${n_cases + 1},\"Automated\")"),
    ("Open defects", f'=COUNTIF(Defects!$K$2:$K${len(defects) + 1},"Open")'),
    ("Fixed defects", f'=COUNTIF(Defects!$K$2:$K${len(defects) + 1},"Fixed")'),
]
for i, (k, f) in enumerate(summary, start=r0 + 1):
    wsum.cell(row=i, column=1, value=k).font = BOLD
    cell = wsum.cell(row=i, column=2, value=f)
    cell.font, cell.alignment = BODY_FONT, Alignment(horizontal="left")
    if k == "Pass rate":
        cell.number_format = "0.0%"

r1 = r0 + len(summary) + 2
wsum.cell(row=r1, column=1, value="By Module").font = Font(name=FONT, size=12, bold=True, color="1F4E78")
mods = list(dict.fromkeys(c["module"] for c in CASES))
for c, h in enumerate(["Module", "", "Total", "Passed"], start=1):
    if h:
        cell = wsum.cell(row=r1 + 1, column=c, value=h)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
for i, m in enumerate(mods, start=r1 + 2):
    wsum.cell(row=i, column=1, value=m).font = BODY_FONT
    mr = f"'Test Cases'!$B$2:$B${n_cases + 1}"
    wsum.cell(row=i, column=3, value=f"=COUNTIF({mr},A{i})").font = BODY_FONT
    wsum.cell(row=i, column=4, value=f'=COUNTIFS({mr},A{i},{st},"Pass")').font = BODY_FONT

r2 = r1 + len(mods) + 3
wsum.cell(row=r2, column=1, value="How to use this document").font = Font(name=FONT, size=12, bold=True,
                                                                          color="1F4E78")
legend = [
    ("Re-executing", "Run the automation suite, then update 'Actual Result', 'Status', 'Executed By' and "
                     "'Execution Date' on the Test Cases sheet. Summary counts update automatically."),
    ("Status values", "Pass / Fail / Blocked / Not Run (drop-down list)"),
    ("Priority", "High = core business rule or data correctness; Medium = robustness; Low = cosmetic/docs"),
    ("Known issues", "Failed cases are linked to the Defects sheet via 'Remarks / Defect ID'"),
    ("Test Data sheet", "Expected (hand-verified) vs Actual (extracted) values per PDF; Result is a formula"),
]
for i, (k, v) in enumerate(legend, start=r2 + 1):
    wsum.cell(row=i, column=1, value=k).font = BOLD
    wsum.cell(row=i, column=2, value=v).font = BODY_FONT
    wsum.cell(row=i, column=2).alignment = WRAP

for sheet in wb.worksheets:
    sheet.sheet_view.zoomScale = 90

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print("Saved", OUT, "-", n_cases, "test cases")
print(recalculate(OUT))
