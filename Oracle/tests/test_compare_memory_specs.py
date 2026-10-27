"""
Tests for compare_memory_specs.py.

Run from the project folder:
    python -m unittest discover -s tests -v

Test groups
-----------
TestHelpers            - small parsing/formatting functions
TestLocators           - the regex "locators" on small hand-written PDF-like text
TestComparisonRules    - Match / Mismatch rules, using fake PDF findings
TestRealPdfs           - golden values extracted from the 26 real PDFs (verified by hand)
TestCsvAgainstPdfs     - the real CSV must fully match the PDFs
TestCatchesWrongData   - negative tests: wrong CSV values / missing columns must be reported
TestEndToEnd           - runs the script like a user and checks the report files and error messages
TestReadme             - README code walkthrough line numbers still point at the right code
TestDependencies       - requirements.txt matches the installed library versions

No existing project file is modified: the end-to-end test writes into a temporary folder.
"""

import csv
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

PROJECT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_DIR))

import compare_memory_specs as cms  # noqa: E402

SCRIPT = PROJECT_DIR / "compare_memory_specs.py"
PDF_DIR = PROJECT_DIR / "Oracle pdfs"
CSV_CANDIDATES = [
    PROJECT_DIR / "oracle_host_memory_specification.csv",
    PROJECT_DIR / "oracle_host_memory_specification (1).csv",
]
CSV_PATH = next((p for p in CSV_CANDIDATES if p.exists()), None)
PDF_INDEX = {p.name.lower(): p for p in PDF_DIR.rglob("*.pdf")}

needs_csv = unittest.skipIf(CSV_PATH is None, "input CSV not found")
needs_pdfs = unittest.skipIf(not PDF_INDEX, "PDF folder not found")


def findings(*values, kind=""):
    return [cms.Finding(v, 1, "test snippet", kind) for v in values]


def specs_from_text(text):
    """Run every extractor over one page of hand-written text, like a mini PDF."""
    specs = cms.PdfSpecs(pages=[cms.normalise(text)])
    page = specs.pages[0]
    cms.extract_slots(page, 1, specs)
    cms.extract_processors(page, 1, specs)
    cms.extract_memory(page, 1, specs)
    cms.extract_dimms(page, 1, specs)
    return specs


def read_csv_rows():
    with CSV_PATH.open(newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def csv_row(file_name):
    return next(dict(r) for r in read_csv_rows() if r["file_name"] == file_name)


# --------------------------------------------------------------------------- helpers

class TestHelpers(unittest.TestCase):

    def test_parse_csv_cell(self):
        self.assertEqual(cms.parse_csv_cell("['24']"), ["24"])
        self.assertEqual(cms.parse_csv_cell("['1.5 TB']"), ["1.5 TB"])
        self.assertEqual(cms.parse_csv_cell("['16 GB', '32 GB']"), ["16 GB", "32 GB"])
        self.assertEqual(cms.parse_csv_cell("24"), ["24"])

    def test_parse_csv_cell_treats_missing_markers_as_empty(self):
        for raw in ("['nan']", "nan", "", None, "['none']", "n/a", "-"):
            with self.subTest(raw=raw):
                self.assertEqual(cms.parse_csv_cell(raw), [])

    def test_parse_csv_size(self):
        self.assertEqual(cms.parse_csv_size("768 GB"), 768)
        self.assertEqual(cms.parse_csv_size("1.5 TB"), 1536)
        self.assertEqual(cms.parse_csv_size("6 tb"), 6144)

    def test_parse_csv_size_rejects_bad_text(self):
        with self.assertRaises(ValueError):
            cms.parse_csv_size("lots of memory")

    def test_word_to_int(self):
        self.assertEqual(cms.word_to_int("Twenty-four"), 24)
        self.assertEqual(cms.word_to_int("Ninety-six"), 96)
        self.assertEqual(cms.word_to_int("sixteen"), 16)
        self.assertEqual(cms.word_to_int("192"), 192)

    def test_format_size(self):
        self.assertEqual(cms.format_size(768), "768 GB")
        self.assertEqual(cms.format_size(1536), "1.5 TB")
        self.assertEqual(cms.format_size(None), "")


# --------------------------------------------------------------------------- regex "locators"

class TestLocators(unittest.TestCase):

    def test_rdimm_is_not_confused_with_lrdimm(self):
        specs = specs_from_text(
            "• RDIMM options: 16 GB at DDR4-2666 and 32 GB at DDR4-2666 "
            "• LRDIMM option: 64 GB at DDR4-2666"
        )
        self.assertEqual(cms.distinct(specs.dimm["RDIMM"]), [16, 32])
        self.assertEqual(cms.distinct(specs.dimm["LRDIMM"]), [64])

    def test_broken_bullet_and_run_on_sidebar_text(self):
        # Regression: x8-2 uses U+FFFD as its bullet and the LRDIMM line runs into sidebar text.
        specs = specs_from_text(
            "� RDIMM options: 16 GB at DDR4-2666 single rank � "
            "LRDIMM option: 64 GB at DDR4-2666 dual rank Oracle Server X8-2 is the most versatile 1.5 TB"
        )
        self.assertEqual(cms.distinct(specs.dimm["LRDIMM"]), [64])
        self.assertEqual(cms.distinct(specs.dimm["RDIMM"]), [16])

    def test_ddr_speed_is_not_read_as_slot_count(self):
        # Regression: "DDR4-2400 DIMMs" was once read as 2400 slots.
        specs = specs_from_text(
            "16 dual inline memory module (DIMM) slots per processor support half and fully "
            "populated memory configurations using 16, 32, or 64 GB DDR4-2400 DIMMs"
        )
        self.assertEqual(cms.distinct(specs.slots_total), [])
        self.assertEqual(cms.distinct(specs.slots_per_proc), [16])

    def test_dimms_per_channel_is_not_a_slot_count(self):
        specs = specs_from_text("1,333 MT/sec at three DIMMs per channel (3 DPC)")
        self.assertEqual(specs.slots_total + specs.slots_per_proc, [])

    def test_word_numbers_and_per_cpu_module(self):
        specs = specs_from_text(
            "Ninety-six DIMM slots (12 per CPU module) provide up to 6 TB of DDR4 ECC DIMM memory"
        )
        self.assertEqual(cms.distinct(specs.slots_total), [96])
        self.assertEqual(cms.distinct(specs.slots_per_proc), [12])
        self.assertEqual(cms.distinct(specs.memory_spec), [6144])

    def test_spec_table_and_summary_are_kept_apart(self):
        specs = specs_from_text(
            "• Twenty-four DIMM slots with maximum memory of 1.5 TB "
            "• Twenty-four DIMM slots provide up to 768 GB of DDR4 DIMM memory"
        )
        self.assertEqual(cms.distinct(specs.memory_summary), [1536])
        self.assertEqual(cms.distinct(specs.memory_spec), [768])
        self.assertEqual(cms.select_memory_value([768], [1536]), (768, "spec table"))

    def test_processor_count(self):
        specs = specs_from_text("Oracle's SPARC T8-4 server is a resilient, four-processor system")
        self.assertEqual(cms.distinct(specs.max_processors), [4])


# --------------------------------------------------------------------------- comparison rules

class TestComparisonRules(unittest.TestCase):

    def rdimm_specs(self, *sizes):
        specs = cms.PdfSpecs()
        specs.dimm["RDIMM"] = findings(*sizes)
        return specs

    def test_dimm_both_empty_is_match(self):
        self.assertEqual(cms.compare_dimm("RDIMM", [], cms.PdfSpecs()).status, cms.MATCH)

    def test_dimm_largest_pdf_option_is_compared(self):
        result = cms.compare_dimm("RDIMM", ["32 GB"], self.rdimm_specs(16, 32))
        self.assertEqual((result.status, result.pdf_value), (cms.MATCH, "32 GB"))

    def test_dimm_different_value_is_mismatch(self):
        result = cms.compare_dimm("RDIMM", ["16 GB"], self.rdimm_specs(16, 32))
        self.assertEqual(result.status, cms.MISMATCH)
        self.assertIn("PDF options: 16 GB, 32 GB", result.comment)

    def test_dimm_missing_on_one_side_is_mismatch(self):
        self.assertEqual(cms.compare_dimm("RDIMM", [], self.rdimm_specs(32)).status, cms.MISMATCH)
        self.assertEqual(cms.compare_dimm("RDIMM", ["32 GB"], cms.PdfSpecs()).status, cms.MISMATCH)

    def memory_specs(self, spec=(), summary=()):
        specs = cms.PdfSpecs()
        specs.memory_spec = findings(*spec, kind="spec")
        specs.memory_summary = findings(*summary, kind="summary")
        return specs

    def test_memory_spec_table_wins_and_conflict_is_noted(self):
        result = cms.compare_memory(["768 GB"], self.memory_specs(spec=[768], summary=[1536]))
        self.assertEqual(result.status, cms.MATCH)
        self.assertIn("internally inconsistent", result.comment)

    def test_memory_matching_only_the_summary_is_mismatch(self):
        result = cms.compare_memory(["1.5 TB"], self.memory_specs(spec=[768], summary=[1536]))
        self.assertEqual(result.status, cms.MISMATCH)
        self.assertIn("appears elsewhere", result.comment)

    def test_memory_largest_configuration_is_used(self):
        result = cms.compare_memory(["8 TB"], self.memory_specs(spec=[4096, 8192]))
        self.assertEqual(result.status, cms.MATCH)
        self.assertIn("smaller configuration", result.comment)

    def test_memory_missing_in_csv_is_mismatch(self):
        result = cms.compare_memory([], self.memory_specs(spec=[16384]))
        self.assertEqual(result.status, cms.MISMATCH)
        self.assertIn("missing in CSV", result.comment)

    def slot_specs(self, total=(), per_proc=(), processors=()):
        specs = cms.PdfSpecs()
        specs.slots_total = findings(*total, kind="total")
        specs.slots_per_proc = findings(*per_proc, kind="per_processor")
        specs.max_processors = findings(*processors)
        return specs

    def test_slots_per_processor_and_derived_total_both_match(self):
        specs = self.slot_specs(per_proc=[16], processors=[4])
        per_proc = cms.compare_slots(["16"], specs)
        derived = cms.compare_slots(["64"], specs)
        self.assertEqual(per_proc.status, cms.MATCH)
        self.assertIn("64 total", per_proc.pdf_value)
        self.assertEqual(derived.status, cms.MATCH)

    def test_slots_other_number_is_mismatch(self):
        self.assertEqual(cms.compare_slots(["32"], self.slot_specs(per_proc=[16], processors=[4])).status,
                         cms.MISMATCH)
        self.assertEqual(cms.compare_slots(["12"], self.slot_specs(total=[24])).status, cms.MISMATCH)

    def test_overall_status(self):
        ok = cms.FieldResult("", "", cms.MATCH, "", [])
        bad = cms.FieldResult("", "", cms.MISMATCH, "", [])
        self.assertEqual(cms.overall({"a": ok, "b": ok}), cms.MATCH)
        self.assertEqual(cms.overall({"a": ok, "b": bad}), cms.MISMATCH)


# --------------------------------------------------------------------------- golden values from the real PDFs

# Values read from the PDFs by hand.
# (system DIMM slots (explicit or derived), slots per processor, max memory, max LRDIMM, max RDIMM)
EXPECTED_PDF_VALUES = {
    "x5-2-datasheet.pdf": (24, None, "768 GB", "32 GB", "16 GB"),
    "x5-2l-datasheet.pdf": (24, None, "768 GB", "32 GB", "16 GB"),
    "x5-4-datasheet-v1.pdf": (96, 24, "3 TB", "32 GB", "16 GB"),
    "x5-8-datasheet.pdf": (192, 24, "6 TB", "32 GB", "16 GB"),
    "x6-2-datasheet-v4.pdf": (24, None, "768 GB", "64 GB", "32 GB"),
    "x6-2l-datasheet-v3.pdf": (24, None, "768 GB", "64 GB", "32 GB"),
    "x7-2_datasheet-v1.pdf": (24, None, "1.5 TB", "64 GB", ""),
    "x7-2l_datasheet.pdf": (24, None, "1.5 TB", "64 GB", "32 GB"),
    "x7-8-dual4socket-datasheet.pdf": (48, 12, "3 TB", "64 GB", "32 GB"),
    "x7-8_8socket_datasheet.pdf": (96, 12, "6 TB", "64 GB", "32 GB"),
    "x8-2-datasheet.pdf": (24, None, "1.5 TB", "64 GB", "32 GB"),
    "x8-2l-datasheet.pdf": (24, None, "1.5 TB", "64 GB", "32 GB"),
    "x8-8-8socket-datasheet.pdf": (96, 12, "6 TB", "64 GB", "32 GB"),
    "x8-8-dual-4socket-datasheet.pdf": (48, 12, "3 TB", "64 GB", "32 GB"),
    "x9-2-datasheet.pdf": (32, None, "2 TB", "", "64 GB"),
    "x9-2l-datasheet.pdf": (32, None, "2 TB", "", "64 GB"),
    "sparc-t8-1-ds-3864201.pdf": (16, 16, "2 TB", "", ""),
    "sparc-t8-2-ds-3864232.pdf": (32, 16, "4 TB", "", ""),
    "sparc-t8-4-ds-3864254.pdf": (64, 16, "8 TB", "", ""),
    "sparc-m8-8-ds-3864183.pdf": (128, 16, "16 TB", "", ""),
    "sparc-s7-2-s7-2l-ds-3049064.pdf": (16, 8, "1 TB", "", ""),
    "netra-s7-2-datasheet-3045526.pdf": (16, 8, "1 TB", "", ""),
    "fujitsu-m10-1-ds-1924203.pdf": (16, 16, "512 GB", "", ""),
    "fujitsu-m12-1-ds-3745915.pdf": (16, 16, "1 TB", "", ""),
    "fujitsu-m12-ds-3669775.pdf": (48, 24, "3 TB", "", ""),
    "fujitsu-m12s-ds-3629110.pdf": (192, 24, "12 TB", "", ""),
}


def largest(found):
    return cms.format_size(max(cms.distinct(found))) if found else ""


@needs_pdfs
class TestRealPdfs(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.specs = {name: cms.extract_pdf(PDF_INDEX[name.lower()])
                     for name in EXPECTED_PDF_VALUES if name.lower() in PDF_INDEX}

    def test_every_expected_pdf_exists_and_has_text(self):
        for name in EXPECTED_PDF_VALUES:
            with self.subTest(pdf=name):
                self.assertIn(name, self.specs, "PDF file missing")
                self.assertTrue(any(page.strip() for page in self.specs[name].pages))

    def test_extracted_values_match_hand_verified_values(self):
        for name, (slots, per_proc, memory, lrdimm, rdimm) in EXPECTED_PDF_VALUES.items():
            specs = self.specs.get(name)
            if specs is None:
                continue  # reported by the test above
            with self.subTest(pdf=name):
                _, total, pp, _, derived = cms.slot_values(specs)
                self.assertEqual(total if total is not None else derived, slots, "system DIMM slots")
                self.assertEqual(pp, per_proc, "slots per processor")
                mem_gb, _ = cms.select_memory_value(cms.distinct(specs.memory_spec),
                                                    cms.distinct(specs.memory_summary))
                self.assertEqual(cms.format_size(mem_gb), memory, "maximum memory")
                self.assertEqual(largest(specs.dimm["LRDIMM"]), lrdimm, "LRDIMM")
                self.assertEqual(largest(specs.dimm["RDIMM"]), rdimm, "RDIMM")
                self.assertEqual(largest(specs.dimm["UDIMM"]), "", "UDIMM")
                self.assertEqual(largest(specs.dimm["SODIMM"]), "", "SODIMM")


# --------------------------------------------------------------------------- the real CSV

@needs_csv
@needs_pdfs
class TestCsvAgainstPdfs(unittest.TestCase):

    def test_csv_has_all_required_columns(self):
        with CSV_PATH.open(newline="", encoding="utf-8-sig") as fh:
            header = next(csv.reader(fh))
        for column in cms.FIELDS + ["file_name"]:
            with self.subTest(column=column):
                self.assertIn(column, header)

    def test_every_csv_row_has_a_pdf(self):
        for row in read_csv_rows():
            with self.subTest(pdf=row["file_name"]):
                self.assertIn(row["file_name"].strip().lower(), PDF_INDEX)

    def test_every_csv_row_matches_its_pdf(self):
        for row in read_csv_rows():
            name = row["file_name"]
            with self.subTest(pdf=name):
                results, error = cms.compare_row(row, PDF_INDEX.get(name.lower()))
                self.assertEqual(error, "")
                problems = {f: (r.csv_value, r.pdf_value, r.comment)
                            for f, r in results.items() if r.status != cms.MATCH}
                self.assertEqual(problems, {})


# --------------------------------------------------------------------------- negative tests

@needs_csv
@needs_pdfs
class TestCatchesWrongData(unittest.TestCase):
    """A checker that always says 'Match' is useless - wrong CSV values must be caught."""

    WRONG_VALUES = [
        ("x8-2-datasheet.pdf", "maximum_RDIMM", "['16 GB']"),        # smaller RDIMM option, not the max
        ("x8-2-datasheet.pdf", "maximum_LRDIMM", "['32 GB']"),
        ("x5-2-datasheet.pdf", "dimm_slots", "['12']"),
        ("sparc-t8-4-ds-3864254.pdf", "dimm_slots", "['32']"),        # neither per-processor nor total
        ("x6-2-datasheet-v4.pdf", "maximum_memory", "['1.5 TB']"),    # summary value, not spec table
        ("sparc-m8-8-ds-3864183.pdf", "maximum_memory", "['nan']"),   # missing in CSV
        ("x7-2_datasheet-v1.pdf", "maximum_RDIMM", "['32 GB']"),      # PDF never mentions RDIMM
        ("x9-2-datasheet.pdf", "maximum_UDIMM", "['16 GB']"),
    ]

    def test_wrong_values_are_reported_as_mismatch(self):
        for name, field_name, wrong_value in self.WRONG_VALUES:
            with self.subTest(pdf=name, field=field_name, value=wrong_value):
                row = csv_row(name)
                row[field_name] = wrong_value
                results, error = cms.compare_row(row, PDF_INDEX[name.lower()])
                self.assertEqual(error, "")
                self.assertEqual(results[field_name].status, cms.MISMATCH)
                self.assertEqual(cms.overall(results), cms.MISMATCH)

    def test_missing_pdf_is_an_error(self):
        results, error = cms.compare_row(csv_row("x5-2-datasheet.pdf"), None)
        self.assertIsNone(results)
        self.assertIn("not found", error)

    def test_bad_size_text_becomes_error_row_not_crash(self):
        row = csv_row("x5-2-datasheet.pdf")
        row["maximum_memory"] = "['a lot']"
        _, results, error = cms.safe_compare_row(row, PDF_INDEX)
        self.assertIsNone(results)
        self.assertIn("Unrecognised size", error)

    def test_missing_csv_column_is_an_error_not_a_match(self):
        # DEF-01: a missing column used to be treated as "empty" and reported as Match.
        row = csv_row("x5-2-datasheet.pdf")
        del row["maximum_UDIMM"]
        results, error = cms.compare_row(row, PDF_INDEX["x5-2-datasheet.pdf"])
        self.assertIsNone(results)
        self.assertIn("missing column(s): maximum_UDIMM", error)


# --------------------------------------------------------------------------- end to end

@needs_csv
@needs_pdfs
class TestEndToEnd(unittest.TestCase):

    def run_script(self, *args):
        return subprocess.run([sys.executable, str(SCRIPT), *args],
                              capture_output=True, text=True, timeout=300)

    def test_script_produces_reports(self):
        with tempfile.TemporaryDirectory() as out_dir:
            run = self.run_script("--csv", str(CSV_PATH), "--out-dir", out_dir)
            self.assertEqual(run.returncode, 0, run.stderr)
            rows = len(read_csv_rows())
            self.assertIn(f"Processed {rows} rows ({rows} match, 0 mismatch", run.stdout)

            with (Path(out_dir) / "memory_spec_comparison.csv").open(encoding="utf-8-sig") as fh:
                report = list(csv.DictReader(fh))
            self.assertEqual(len(report), rows)
            self.assertTrue(all(r["overall_status"] == cms.MATCH for r in report))

            workbook = load_workbook(Path(out_dir) / "memory_spec_comparison.xlsx")
            self.assertEqual(workbook.sheetnames, ["Results", "Summary", "Evidence"])
            self.assertEqual(workbook["Results"].max_row, rows + 1)
            summary_header = [c.value for c in workbook["Summary"][1]]
            self.assertEqual(summary_header, ["field", cms.MATCH, cms.MISMATCH])  # DEF-04: no unused Review column
            self.assertGreater(workbook["Evidence"].max_row, rows)
            workbook.close()

    def test_default_run_without_options(self):
        # DEF-02: README usage "python compare_memory_specs.py" must work with the downloaded CSV name.
        with tempfile.TemporaryDirectory() as out_dir:
            run = self.run_script("--out-dir", out_dir)
            self.assertEqual(run.returncode, 0, run.stderr)
            self.assertIn(f"Input: {CSV_PATH}", run.stdout)

    def test_missing_csv_file_fails_with_message(self):
        with tempfile.TemporaryDirectory() as out_dir:
            run = self.run_script("--csv", str(Path(out_dir) / "nope.csv"), "--out-dir", out_dir)
            self.assertEqual(run.returncode, 2)
            self.assertIn("input CSV not found", run.stderr)
            self.assertNotIn("Traceback", run.stderr)

    def test_missing_pdf_folder_fails_with_message(self):
        with tempfile.TemporaryDirectory() as out_dir:
            run = self.run_script("--csv", str(CSV_PATH), "--pdf-dir", str(Path(out_dir) / "none"),
                                  "--out-dir", out_dir)
            self.assertEqual(run.returncode, 2)
            self.assertIn("PDF folder not found", run.stderr)

    def test_csv_file_missing_a_column_fails_with_message(self):
        # DEF-01 at file level: the whole run is rejected before any comparison.
        with tempfile.TemporaryDirectory() as out_dir:
            bad_csv = Path(out_dir) / "no_udimm.csv"
            with CSV_PATH.open(newline="", encoding="utf-8-sig") as src, \
                    bad_csv.open("w", newline="", encoding="utf-8") as dst:
                reader = csv.DictReader(src)
                columns = [c for c in reader.fieldnames if c != "maximum_UDIMM"]
                writer = csv.DictWriter(dst, fieldnames=columns, extrasaction="ignore")
                writer.writeheader()
                writer.writerows(reader)
            run = self.run_script("--csv", str(bad_csv), "--out-dir", out_dir)
            self.assertEqual(run.returncode, 2)
            self.assertIn("missing required column(s): maximum_UDIMM", run.stderr)
            self.assertFalse((Path(out_dir) / "memory_spec_comparison.csv").exists())


class TestReadme(unittest.TestCase):
    """DEF-03: keep the README code walkthrough in sync with the script."""

    ROW = re.compile(r"^\| (\d+)-(\d+) \| `(\w+)` \|", re.M)

    def test_walkthrough_line_numbers_match_the_code(self):
        readme = (PROJECT_DIR / "README.md").read_text(encoding="utf-8")
        code = SCRIPT.read_text(encoding="utf-8").splitlines()
        rows = self.ROW.findall(readme)
        self.assertGreater(len(rows), 50, "walkthrough table not found")
        for start, end, name in rows:
            with self.subTest(name=name, lines=f"{start}-{end}"):
                first = code[int(start) - 1]
                self.assertRegex(first, rf"^(def |class )?{name}\b", f"line {start} is: {first!r}")
                self.assertLessEqual(int(end), len(code))

    def test_every_top_level_definition_is_documented(self):
        readme = (PROJECT_DIR / "README.md").read_text(encoding="utf-8")
        documented = {name for _, _, name in self.ROW.findall(readme)}
        defined = set(re.findall(r"^(?:def |class )?([A-Za-z_]\w*)(?:\(| =|:)", SCRIPT.read_text(encoding="utf-8"), re.M))
        defined -= {"if", "import", "from", "MISMATCH"}
        self.assertEqual(sorted(defined - documented), [])


class TestDependencies(unittest.TestCase):
    """DEF-05: requirements.txt pins the versions the tests were run with."""

    def test_requirements_match_installed_versions(self):
        from importlib.metadata import version
        pins = dict(line.strip().split("==") for line in
                    (PROJECT_DIR / "requirements.txt").read_text(encoding="utf-8").splitlines() if "==" in line)
        self.assertEqual(set(pins), {"pymupdf", "openpyxl"})
        for package, pinned in pins.items():
            with self.subTest(package=package):
                self.assertEqual(version(package), pinned)


if __name__ == "__main__":
    unittest.main()
